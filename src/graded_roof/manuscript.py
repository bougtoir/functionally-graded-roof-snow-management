from __future__ import annotations

import csv
import hashlib
import json
import re
import zipfile
from copy import deepcopy
from pathlib import Path

import pandas as pd
import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt

TITLE = (
    "Functionally graded roofs for passive snow management: "
    "nominal trade-off expansion limited by constructability and timestep sensitivity"
)

FIGURE_CAPTIONS = [
    "Figure 1. Reduced-order model and comparative optimization workflow.",
    "Figure 2. Synthetic snowfall scenarios and JMA ground-snow contexts.",
    "Figure 3. Unique objective pairs for the exhaustive uniform and heuristic joint "
    "sets, shown over the full range and the intermediate trade-off region.",
    "Figure 4. Continuous joint-knee slope and friction profiles compared with the "
    "post hoc rule-compliant segmented, generic-class mapping; the mapped profile is "
    "not a constrained optimum.",
    "Figure 5. Unique objective pairs for the uniform, geometry-only, surface-only, "
    "and joint sets.",
    "Figure 6. Four-metric numerical-resolution comparison relative to the 48-cell, 0.5-h "
    "reference; the dashed line is the frozen relative-error tolerance.",
    "Figure 7. Prespecified one-at-a-time sensitivity. Labels state each low/high "
    "factor or encoded alternative; bars are relative changes from baseline.",
    "Figure 8. Candidate-specific Q95 retained-mass and release-mass outcomes under "
    "250 paired perturbation draws. Symmetric-log axes include zero outcomes.",
    "Figure 9. Frozen labor-scarcity selections in the base decision scenario and "
    "phase-diagram objective pairs; bubble area reflects selection frequency.",
]

FIGURE_FILENAMES = [
    "figure_1_conceptual_model.png",
    "figure_2_weather_scenarios.png",
    "figure_3_primary_pareto_comparison.png",
    "figure_4_selected_graded_profile.png",
    "figure_5_ablation_pareto_fronts.png",
    "figure_6_numerical_convergence.png",
    "figure_7_sensitivity.png",
    "figure_8_monte_carlo_robustness.png",
    "figure_9_decision_phase_diagram.png",
]


def _format_number(value: float) -> str:
    if abs(value) >= 100:
        return f"{value:.1f}"
    if abs(value) >= 10:
        return f"{value:.2f}"
    return f"{value:.3f}"


def _add_line_numbering(document: Document) -> None:
    for section in document.sections:
        line_number = OxmlElement("w:lnNumType")
        line_number.set(qn("w:countBy"), "1")
        line_number.set(qn("w:start"), "1")
        line_number.set(qn("w:restart"), "continuous")
        section._sectPr.append(line_number)


def _configure_document(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Arial"
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.start_type = WD_SECTION.CONTINUOUS
    _add_line_numbering(document)


def _add_equation(document: Document, equation: str) -> None:
    paragraph = document.add_paragraph()
    math_paragraph = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = equation
    run.append(text)
    math.append(run)
    math_paragraph.append(math)
    paragraph._p.append(math_paragraph)


def _enforce_ascii(document: Document) -> None:
    replacements = {
        "\u2013": "-",
        "\u2014": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2212": "-",
        "\u00d7": "x",
        "\u2264": "<=",
        "\u2022": "-",
    }
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            text = run.text
            for original, replacement in replacements.items():
                text = text.replace(original, replacement)
            text.encode("ascii")
            if text != run.text:
                run.text = text
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        text = run.text
                        for original, replacement in replacements.items():
                            text = text.replace(original, replacement)
                        text.encode("ascii")
                        if text != run.text:
                            run.text = text
    _apply_unit_superscripts(document)


def _apply_unit_superscripts(document: Document) -> None:
    unit_exponent = re.compile(r"\b([mhs])(-[123])\b")

    def transform_paragraph(paragraph) -> None:
        for run in list(paragraph.runs):
            matches = list(unit_exponent.finditer(run.text))
            if not matches:
                continue
            parent = run._r.getparent()
            insertion_index = parent.index(run._r)
            run_properties = run._r.find(qn("w:rPr"))
            cursor = 0
            pieces: list[tuple[str, bool]] = []
            for match in matches:
                pieces.append((run.text[cursor : match.start(2)], False))
                pieces.append((match.group(2).replace("-", "−"), True))
                cursor = match.end(2)
            pieces.append((run.text[cursor:], False))
            for text, superscript in pieces:
                if not text:
                    continue
                new_run = OxmlElement("w:r")
                if run_properties is not None:
                    new_run.append(deepcopy(run_properties))
                if superscript:
                    properties = new_run.find(qn("w:rPr"))
                    if properties is None:
                        properties = OxmlElement("w:rPr")
                        new_run.insert(0, properties)
                    vertical_alignment = OxmlElement("w:vertAlign")
                    vertical_alignment.set(qn("w:val"), "superscript")
                    properties.append(vertical_alignment)
                new_text = OxmlElement("w:t")
                if text[:1].isspace() or text[-1:].isspace():
                    new_text.set(qn("xml:space"), "preserve")
                new_text.text = text
                new_run.append(new_text)
                parent.insert(insertion_index, new_run)
                insertion_index += 1
            parent.remove(run._r)

    for paragraph in document.paragraphs:
        transform_paragraph(paragraph)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    transform_paragraph(paragraph)


def _add_title_page(document: Document, metadata: dict) -> None:
    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(TITLE)
    for author in metadata["authors"]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(author["name"])
        if author.get("corresponding"):
            paragraph.add_run("*")
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(author["affiliation"])
    corresponding = next(
        author for author in metadata["authors"] if author.get("corresponding")
    )
    document.add_paragraph(
        "Corresponding author: "
        f"{corresponding['name']}, {corresponding['postal_address']}, "
        f"{corresponding['email']}"
    )


def _add_keywords(document: Document) -> None:
    document.add_paragraph(
        "Keywords: roof snow; passive snow management; multiobjective optimization; "
        "spatial heterogeneity; snow shedding; reduced-order model"
    )


def _add_dataframe_table(
    document: Document,
    frame: pd.DataFrame,
    caption: str,
    columns: list[str],
    *,
    maximum_rows: int | None = None,
    headers: list[str] | None = None,
) -> None:
    document.add_paragraph(caption, style="Heading 3")
    shown = frame.loc[:, columns]
    if maximum_rows is not None:
        shown = shown.head(maximum_rows)
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for index, column in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.text = (
            headers[index] if headers is not None else column
        )
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(7)
    for row in shown.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if isinstance(value, float):
                cells[index].text = _format_number(value)
            else:
                cells[index].text = str(value)
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(7)


def _reference_text(row: pd.Series) -> str:
    author_entries = [
        entry.strip() for entry in str(row["authors"]).split(";")
    ]
    formatted_authors: list[str] = []
    for entry in author_entries:
        if " et al." in entry:
            parts = entry.split()
            initials = ".".join(parts[1]) + "."
            formatted_authors.append(f"{parts[0]}, {initials}, et al.")
            continue
        parts = entry.split()
        if len(parts) >= 2 and parts[-1].isupper():
            initials = ".".join(parts[-1]) + "."
            formatted_authors.append(
                f"{' '.join(parts[:-1])}, {initials}"
            )
        else:
            formatted_authors.append(entry)
    authors = ", ".join(formatted_authors)
    identifier = str(row["doi_or_identifier"])
    locator = (
        f"https://doi.org/{identifier}."
        if identifier.startswith("10.")
        else f"Available at: {row['official_url']}."
    )
    volume = str(row.get("volume", "")).removesuffix(".0")
    issue = str(row.get("issue", "")).removesuffix(".0")
    pages = str(row.get("pages_or_article_number", "")).removesuffix(".0")
    bibliographic = str(row["journal_or_publisher"])
    if volume and volume.lower() != "nan":
        bibliographic += f" {volume}"
        if issue and issue.lower() != "nan":
            bibliographic += f"({issue})"
    if pages and pages.lower() != "nan":
        bibliographic += f", {pages}"
    return (
        f"{authors}, {int(row['year'])}. {row['title']}. "
        f"{bibliographic}. {locator}"
    )


def _citation_author(entry: str) -> str:
    if " et al." in entry:
        return entry.split()[0] + " et al."
    parts = entry.split()
    if len(parts) >= 2 and parts[-1].isupper():
        return " ".join(parts[:-1])
    return entry


def _citation_label(row: pd.Series) -> str:
    authors = [
        _citation_author(entry.strip())
        for entry in str(row["authors"]).split(";")
    ]
    if len(authors) == 1:
        author_text = authors[0]
    elif len(authors) == 2:
        author_text = f"{authors[0]} and {authors[1]}"
    else:
        author_text = f"{authors[0]} et al."
    return f"{author_text}, {int(row['year'])}"


def _citation_group(references: pd.DataFrame, record_ids: list[str]) -> str:
    selected = references.loc[references["record_id"].isin(record_ids)].copy()
    selected["citation_label"] = selected.apply(_citation_label, axis=1)
    selected["sort_author"] = selected["authors"].str.casefold()
    selected = selected.sort_values(["sort_author", "year", "record_id"])
    return "(" + "; ".join(selected["citation_label"]) + ")"


def _frontier_result_sentence(
    frontier_summary: pd.DataFrame,
) -> str:
    uniform = frontier_summary.set_index("design_class").loc["uniform"]
    joint = frontier_summary.set_index("design_class").loc["joint"]
    return (
        f"The uniform set contained {int(uniform['unique_objective_pairs'])} "
        "unique objective pairs, whereas the joint set contained "
        f"{int(joint['unique_objective_pairs'])}, including the uniform endpoints "
        "and additional intermediate trade-offs. Under the common audit reference, "
        "their normalized hypervolumes were "
        f"{_format_number(uniform['normalized_hypervolume'])} and "
        f"{_format_number(joint['normalized_hypervolume'])}; joint additive epsilon "
        f"versus uniform was {_format_number(joint['additive_epsilon_vs_uniform'])}."
    )


def build_manuscript(root: Path) -> Path:
    results = root / "results" / "generated"
    tables = root / "tables" / "generated"
    metadata = yaml.safe_load(
        (root / "manuscript" / "author_metadata.yaml").read_text(encoding="utf-8")
    )
    config = yaml.safe_load(
        (root / "config" / "production.yaml").read_text(encoding="utf-8")
    )
    baseline = pd.read_csv(results / "baseline_aggregate.csv")
    convergence = pd.read_csv(results / "convergence.csv")
    frontier_summary = pd.read_csv(
        results / "final_revision_frontier_summary.csv"
    )
    frontier_hypervolume = pd.read_csv(
        results / "final_revision_frontier_metric_sensitivity.csv"
    )
    frontier_hypervolume = frontier_hypervolume.loc[
        frontier_hypervolume["normalization_origin"].eq("zero")
        & frontier_hypervolume["reference_margin"].eq(1.05),
        ["design_class", "normalized_hypervolume_fraction"],
    ].rename(
        columns={
            "normalized_hypervolume_fraction": "normalized_hypervolume"
        }
    )
    frontier_summary = frontier_summary.merge(
        frontier_hypervolume, on="design_class", validate="one_to_one"
    ).rename(
        columns={
            "additive_epsilon_vs_uniform_common_minmax": (
                "additive_epsilon_vs_uniform"
            )
        }
    )
    knee_audit = pd.read_csv(results / "final_revision_knee_audit.csv")
    unique_regions = pd.read_csv(
        results / "final_revision_joint_unique_regions.csv"
    )
    robust_equivalent = pd.read_csv(
        results / "final_revision_objective_equivalent_robustness.csv"
    )
    robust_ties = pd.read_csv(
        results / "final_revision_robust_selection_ties.csv"
    )
    discretization_loss = pd.read_csv(
        results / "final_revision_discretization_loss.csv"
    ).iloc[0]
    constructability = pd.read_csv(
        results / "final_revision_constructability_by_objective.csv"
    )
    timestep_audit = pd.read_csv(
        results / "final_revision_selected_design_timestep_audit.csv"
    )
    quarter_hour_audit = pd.read_csv(
        results / "final_revision_selected_design_0p25h_audit.csv"
    )
    constraints = pd.read_csv(
        results / "final_revision_prespecified_constraints.csv"
    )
    jma_pairs = pd.read_csv(
        results / "final_revision_jma_paired_tradeoffs.csv"
    )
    jma_summary = pd.read_csv(
        results / "final_revision_jma_paired_summary.csv"
    ).set_index("metric")
    table_1 = pd.read_csv(tables / "table_1_pareto_summary.csv")
    table_2 = pd.read_csv(tables / "table_2_selected_profile.csv")
    table_3 = pd.read_csv(tables / "table_3_discretization.csv")
    table_5 = pd.read_csv(tables / "table_5_labor_scarcity.csv")
    table_6 = pd.read_csv(tables / "table_6_complexity_analysis.csv")
    table_7 = pd.read_csv(tables / "table_7_complexity_penalty.csv")
    table_8 = pd.read_csv(tables / "table_8_robustness_summary.csv")
    references = pd.read_csv(
        root / "references" / "final_revision_reference_audit.csv"
    )
    roof_model_citations = _citation_group(
        references,
        [f"CRST{index:02d}" for index in range(1, 7)],
    )
    adjacent_study_citations = _citation_group(
        references,
        [f"CRST{index:02d}" for index in range(7, 13)],
    )
    japanese_citations = _citation_group(
        references,
        [f"JP{index:02d}" for index in range(1, 11)],
    )
    accepted = convergence[convergence["within_tolerance"]]
    convergence_text = (
        f"The frozen {100 * config['convergence']['relative_tolerance']:.0f}% "
        "four-metric convergence tolerance was met by "
        f"{len(accepted)} of {len(convergence)} tested resolution combinations."
    )
    one_hour_timestep = timestep_audit.loc[
        timestep_audit["dt_hours"].eq(1.0)
    ]
    production_against_quarter_hour = quarter_hour_audit.loc[
        quarter_hour_audit["dt_hours"].eq(config["simulation"]["dt_hours"])
    ]
    production_step_text = f"{config['simulation']['dt_hours']:g}"
    failed_one_hour_count = int(
        (~one_hour_timestep["within_tolerance"]).sum()
    )
    failed_quarter_hour_count = int(
        (~production_against_quarter_hour["within_frozen_tolerance"]).sum()
    )
    uniform_pairs = unique_regions.loc[
        unique_regions["attainable_by_uniform_under_same_caps"]
    ]
    uniform_pair_text = " and ".join(
        f"{_format_number(row.l_max_kg_per_m)}/"
        f"{_format_number(row.s_max_kg_per_m)} kg m-1"
        for row in uniform_pairs.itertuples(index=False)
    )
    joint_unique_count = int(
        (~unique_regions["attainable_by_uniform_under_same_caps"]).sum()
    )
    joint_design_rows = int(
        frontier_summary.set_index("design_class").loc[
            "joint", "nondominated_design_rows"
        ]
    )
    uniform_unique_count = int(
        frontier_summary.set_index("design_class").loc[
            "uniform", "unique_objective_pairs"
        ]
    )
    retention = baseline.loc[
        baseline["design"] == "retention_non_shedding"
    ].iloc[0]
    shedding = baseline.loc[
        baseline["design"] == "conventional_shedding"
    ].iloc[0]

    document = Document()
    _configure_document(document)
    _add_title_page(document, metadata)
    document.add_heading("Abstract", level=1)
    abstract = (
        "Passive roof-snow strategies trade retained modeled snow mass against "
        "discrete eave-release mass. We tested whether spatial variation in slope, "
        "friction, and adhesion changed that trade-off relative to an exhaustively "
        "mapped uniform design space. A unit-width, cell-based threshold-release model "
        "was evaluated under synthetic weather, with Japanese Meteorological Agency "
        "daily observations used only as supplementary forcing. Multi-seed NSGA-II "
        "generated geometry-only, surface-only, and joint heterogeneous sets. "
        + _frontier_result_sentence(frontier_summary)
        + f" The joint set had {joint_unique_count} "
        "intermediate objective pairs unavailable to the uniform set under the same "
        f"objective caps. {len(robust_ties)} candidates tied the minimum robustness "
        "score. Mapping the continuous joint knee to rule-compliant generic segments "
        "changed modeled retained mass by "
        f"{discretization_loss['l_max_percent_change']:+.2f}% and release mass by "
        f"{discretization_loss['s_max_percent_change']:+.2f}%, shifting the selected "
        "point toward the low-retention, high-release shedding regime. "
        f"A targeted 0.25-h sensitivity failed the four-metric tolerance for "
        f"{failed_quarter_hour_count} of {len(production_against_quarter_hour)} selected "
        f"designs relative to the {production_step_text}-h production results. Spatial "
        "grading produced nominal intermediate trade-offs absent from the exhaustive "
        "uniform grid, but constructability and interval sensitivity substantially "
        "limit interpretation; no front-wide "
        "superiority, structural safety, or validated roof performance is established."
    )
    document.add_paragraph(abstract)
    _add_keywords(document)

    document.add_heading("1. Introduction", level=1)
    document.add_paragraph(
        "Roof snow is spatially variable, evolves with weather, and can either remain "
        "on the roof or leave in discrete releases. Roof-scale experiments and "
        "deposition studies show that geometry, wind, and scale influence snow "
        "distributions, while validation remains essential for computational models "
        f"{roof_model_citations}. Adjacent photovoltaic, membrane-roof, "
        "meteorological, and field studies "
        "further motivate explicit sensitivity analysis and independent validation "
        f"{adjacent_study_citations}."
    )
    document.add_paragraph(
        "Japanese technical and administrative sources distinguish retention, "
        "shedding, melting, and load-resisting strategies and document the importance "
        "of surface condition, aging, snow guards, drainage, cornices, and site context "
        f"{japanese_citations}. Retention can reduce required snow-fall space at dense "
        "sites, whereas "
        "shedding requires a suitable receiving zone; retained snow instead imposes "
        "long-duration load and drainage obligations. These sources do not establish "
        "universal material coefficients, safety thresholds, or one preferable strategy. "
        "They motivate a conditional comparison of generic uniform and graded design "
        "spaces."
    )
    document.add_paragraph(
        "The research question was frozen before production analysis: under identical "
        "modeled environmental conditions, how does allowing spatial variation in roof "
        "geometry and snow-surface interaction change the achievable trade-off between "
        "retained modeled snow mass and discrete shedding events compared with optimized "
        "uniform roofs? The hypothesis that heterogeneity expands the attainable design "
        "space was tested without assuming superiority."
    )

    document.add_heading("2. Methods", level=1)
    document.add_heading("2.1 Reduced-order roof-snow model", level=2)
    document.add_paragraph(
        f"An {config['roof']['length_m']:g} m long, "
        f"{config['roof']['width_m']:g} m wide roof was represented by "
        f"{config['roof']['cells']} ridge-to-eave cells. Each "
        "cell carried slope, static friction, kinetic friction, adhesion, snow mass, "
        "mass-weighted age, volume, and density. Snowfall was added before each "
        "within-step maximum was recorded. Melt removed mass, volume, and age moment "
        "proportionally. Compaction changed volume and diagnostic depth. Density "
        "converted observed snowfall depth to mass but was not asserted to change "
        "basal sliding at fixed mass."
    )
    document.add_paragraph(
        "A cell became mobile when downslope gravity exceeded static friction plus "
        "area-scaled adhesion. Kinetic friction controlled transport after release. "
        "Mass, age moment, and volume moved together. Snow leaving the final cell was "
        "an eave-shedding event. The kinetic-energy metric was a roof-wide specific-"
        "energy proxy, not an impact or injury model."
    )
    _add_equation(document, "F_g = m g sin(theta)")
    _add_equation(document, "F_r = mu_s m g cos(theta) + A tau_a")
    document.add_heading("2.2 Outcomes and comparison sets", level=2)
    document.add_paragraph(
        "The locked primary outcomes were Lmax, the maximum modeled roof snow mass "
        "per metre width immediately after snowfall and before same-step melt or "
        "release, and Smax, the maximum eave-release mass in one timestep per metre "
        "width. Secondary outcomes included total shed mass, event count, event-mass "
        "summaries, Snow Shedding Concentration Index (SSCI), kinetic-energy proxy, "
        "intervention-threshold time, manual triggers, residual mass, maximum modeled "
        "snow depth, total variation, transition count, and discretization loss."
    )
    _add_equation(document, "SSCI = sum_j (m_j / M_total)^2")
    document.add_paragraph(
        "Comparators were retention-oriented, shedding-oriented, and intermediate "
        "uniform baselines; the complete optimized uniform frontier; and geometry-only, "
        "surface-only, and joint heterogeneous frontiers. Figure 1 summarizes the "
        "model and optimization workflow; Figure 2 summarizes synthetic and observed "
        "weather contexts; Figures 3–5 show the primary frontier, selected profile, "
        "and ablation frontiers. Tables 1–3 report frontier metrics, the selected "
        "profile, and continuous-to-discrete mapping."
    )
    document.add_heading("2.3 Weather forcing", level=2)
    document.add_paragraph(
        f"Primary optimization used "
        f"{len(config['synthetic_weather']['scenarios'])} deterministic synthetic "
        "winters spanning cold "
        "persistent, maritime warming, and severe inland conditions, including repeated "
        "snowfall and configured rain-on-snow periods. JMA daily mean temperature, "
        "precipitation, snowfall depth, and ground snow depth were downloaded as "
        "immutable official HTML snapshots for "
        f"{len(config['jma']['stations'])} stations and "
        f"{len(config['jma']['winters'])} "
        "winters. Original value strings and parsed numeric fields were retained. "
        "JMA ground observations were used only as scenario forcing and were not treated "
        "as direct roof-load measurements."
    )
    document.add_heading("2.4 Optimization and analysis", level=2)
    document.add_paragraph(
        f"A full {config['optimizer']['uniform_slope_grid']} x "
        f"{config['optimizer']['uniform_friction_grid']} x "
        f"{config['optimizer']['uniform_adhesion_grid']} uniform "
        "slope-friction-adhesion map preceded "
        "heterogeneous optimization. Geometry-only, surface-only, and joint profiles "
        f"were parameterized by {config['optimizer']['profile_control_points']} "
        "control points interpolated across the roof. Checkpointed NSGA-II used a "
        f"population of {config['optimizer']['population']}, "
        f"{config['optimizer']['generations']} generations, and seeds "
        + ", ".join(str(seed) for seed in config["optimizer"]["seeds"])
        + ". All feasible evaluations from every seed were combined "
        "and filtered to approximate nondominated heterogeneous sets; only the uniform "
        "grid was exhaustively enumerated. Objective pairs were deduplicated at six "
        "decimal places for frontier metrics while design rows remained available for "
        "robustness and constructability analyses. Normalized hypervolume used common "
        "bounds and a reference at 105% of the joint objective maxima. Additive epsilon "
        "versus uniform, combined-set dominance, matched-objective constraints, and "
        "reference-point sensitivity were calculated. A normalized-distance knee was "
        "reported only as a within-front descriptive selection; common and front-specific "
        "normalizations were audited separately. Per-seed metrics summarized optimizer "
        "stochasticity in a machine-readable supplement. "
        f"Production used a {production_step_text}-h timestep, the finest resolution "
        "in the frozen grid. A targeted 0.25-h reevaluation of selected designs was "
        "reported as post-freeze numerical sensitivity rather than a second optimization. "
        "Figure 6 reports the numerical-resolution comparison; Figure 7 reports "
        "prespecified one-at-a-time sensitivity."
    )
    document.add_paragraph(
        f"Monte Carlo robustness used {config['robustness']['samples']} deterministic "
        "draws perturbing friction, adhesion, fresh-snow density, snowfall, and "
        "temperature for prespecified candidates sampled across the uniform and joint "
        "fronts. Every candidate received the same factor vector within each draw, "
        "permitting paired comparisons; Monte Carlo quantiles were not interpreted as "
        "empirical confidence intervals. Continuous optimization enforced only maximum "
        "adjacent slope change. Minimum segment length, maximum transition count, and "
        "generic surface classes were audited after optimization. Continuous joint "
        "profiles were mapped to "
        f"{len(config['surface']['discrete_classes'])} generic "
        "surface classes and piecewise roof segments of at least "
        f"{config['roof']['minimum_segment_cells']} cells, with at most "
        f"{config['roof']['maximum_transitions']} transitions, without claiming "
        "that the mapped design was a constrained optimum or represented proprietary "
        "material performance. Frozen "
        "manual-intervention weights generated low, medium, and high labor-scarcity "
        "choices across configured snowfall, shedding-penalty, and allowable-release "
        "scenarios. Figures 8-9 and Tables 4-8 summarize these analyses."
    )

    document.add_heading("3. Results", level=1)
    document.add_heading("3.1 Baselines and numerical resolution", level=2)
    document.add_paragraph(
        "Across the synthetic scenarios, the retention baseline produced "
        f"Lmax={_format_number(retention['l_max_kg_per_m'])} kg m−1 and "
        f"Smax={_format_number(retention['s_max_kg_per_m'])} kg m−1, whereas "
        "the shedding baseline produced "
        f"Lmax={_format_number(shedding['l_max_kg_per_m'])} kg m−1 and "
        f"Smax={_format_number(shedding['s_max_kg_per_m'])} kg m−1. "
        "These are modeled unit-width masses, not structural loads. "
        + convergence_text
        + f" The {production_step_text}-h interval was the frozen numerical reference "
        f"and production resolution. At 1 h, {failed_one_hour_count} of "
        f"{len(one_hour_timestep)} selected uniform, continuous-joint, and mapped-joint "
        "designs failed the same composite criterion relative to 0.5 h. Comparing the "
        "0.5-h production outcomes against targeted 0.25-h reevaluations, "
        f"{failed_quarter_hour_count} of {len(production_against_quarter_hour)} selected "
        "designs failed the frozen four-metric tolerance. "
        "Accordingly, 0.5 h is described as the production reference resolution, not "
        "as demonstrated convergence beyond that interval. Smax is release mass in one "
        "model interval and therefore must be interpreted with the stated timestep."
    )
    document.add_heading("3.2 Optimized uniform reference", level=2)
    uniform_summary = frontier_summary.set_index("design_class").loc["uniform"]
    document.add_paragraph(
        f"The exhaustive uniform grid produced "
        f"{int(uniform_summary['nondominated_design_rows'])} nondominated design rows "
        "but only "
        f"{int(uniform_summary['unique_objective_pairs'])} objective pairs: "
        f"{uniform_pair_text} for Lmax/Smax. "
        "The multiplicity reflects objective-equivalent profiles rather than a dense "
        "uniform trade-off curve. Surface-only optimization reproduced the same two "
        "objective pairs."
    )
    document.add_heading("3.3 Heterogeneous Pareto comparisons", level=2)
    document.add_paragraph(_frontier_result_sentence(frontier_summary))
    summary_parts = []
    for row in frontier_summary.itertuples(index=False):
        summary_parts.append(
            f"{row.design_class}: {int(row.unique_objective_pairs)} objective pairs, "
            f"normalized hypervolume {_format_number(row.normalized_hypervolume)}, "
            "additive epsilon versus uniform "
            f"{_format_number(row.additive_epsilon_vs_uniform)}"
        )
    document.add_paragraph(
        "The complete frontier summary was " + "; ".join(summary_parts) + ". "
        "Joint hypervolume exceeded uniform because the joint set included the uniform "
        f"objective pairs and {joint_unique_count} additional intermediate pairs. "
        "Joint additive "
        "epsilon versus uniform was nevertheless zero because every uniform reference "
        "point was weakly matched; this does not imply front-wide superiority. "
        "Hypervolume ordering was unchanged under the audited origin and reference-margin "
        "choices."
    )
    joint_knees = knee_audit.loc[
        (knee_audit["design_class"] == "joint")
    ].set_index("normalization")
    smax_limit = constraints.loc[
        constraints["s_max_limit_kg_per_m"].eq(
            constraints["s_max_limit_kg_per_m"].min()
        )
    ].set_index("design_class")
    document.add_paragraph(
        "The front-specific joint knee was "
        f"{_format_number(joint_knees.loc['front_specific_minmax', 'l_max_kg_per_m'])}/"
        f"{_format_number(joint_knees.loc['front_specific_minmax', 's_max_kg_per_m'])} "
        "kg m-1, whereas common normalization selected "
        f"{_format_number(joint_knees.loc['common_combined_minmax', 'l_max_kg_per_m'])}/"
        f"{_format_number(joint_knees.loc['common_combined_minmax', 's_max_kg_per_m'])} "
        "kg m-1. The normalization dependence precludes a cross-front percentage "
        "headline. At the strictest prespecified Smax limit of "
        f"{smax_limit['s_max_limit_kg_per_m'].iloc[0]:.0f} kg m-1, the minimum observed "
        "Lmax was "
        f"{_format_number(smax_limit.loc['joint', 'minimum_l_max_kg_per_m'])} kg m-1 "
        "for joint and "
        f"{_format_number(smax_limit.loc['uniform', 'minimum_l_max_kg_per_m'])} kg m-1 "
        "for uniform."
    )
    document.add_heading("3.4 Robustness", level=2)
    equivalent_80 = robust_equivalent.sort_values(
        "candidate_count", ascending=False
    ).iloc[0]
    equivalent_pair_text = (
        f"{_format_number(equivalent_80['nominal_l_max_kg_per_m'])}/"
        f"{_format_number(equivalent_80['nominal_s_max_kg_per_m'])} kg m-1"
    )
    document.add_paragraph(
        f"Among {int(equivalent_80['candidate_count'])} candidates with the nominal "
        f"{equivalent_pair_text} objective pair, Q95 Lmax ranged from "
        f"{_format_number(equivalent_80['q95_l_min_kg_per_m'])} to "
        f"{_format_number(equivalent_80['q95_l_max_kg_per_m'])} kg m-1, Q95 Smax "
        "ranged from "
        f"{_format_number(equivalent_80['q95_s_min_kg_per_m'])} to "
        f"{_format_number(equivalent_80['q95_s_max_kg_per_m'])} kg m-1, and manual-"
        "intervention probability ranged from "
        f"{equivalent_80['intervention_probability_min']:.3f} to "
        f"{equivalent_80['intervention_probability_max']:.3f}. "
        f"The minimum scalar robustness score was shared by {len(robust_ties)} "
        "candidates, so the retained selected flag is only a deterministic row-order "
        "tie-break. Monte Carlo quantiles describe the configured perturbation model, "
        "not empirical confidence intervals."
    )
    document.add_heading("3.5 Constructability and decision scenarios", level=2)
    feasible_joint = constructability.loc[
        (constructability["design_class"] == "joint")
        & constructability["feasible_design_rows"].gt(0)
    ]
    feasible_joint_count = int(
        constructability.loc[
            constructability["design_class"].eq("joint"),
            "feasible_design_rows",
        ].sum()
    )
    if feasible_joint.empty:
        feasibility_sentence = (
            f"No joint design row among the {joint_design_rows} nondominated rows "
            "met all post hoc constructability checks."
        )
    else:
        feasible_pair_text = (
            f"{_format_number(feasible_joint.iloc[0]['l_max_kg_per_m'])}/"
            f"{_format_number(feasible_joint.iloc[0]['s_max_kg_per_m'])} kg m-1"
        )
        feasibility_sentence = (
            f"Only {feasible_joint_count} of the {joint_design_rows} joint design rows "
            "met all post hoc constructability checks, and all belonged to the "
            f"{feasible_pair_text} objective pair."
        )
    continuous_knee_constructability = constructability.loc[
        (constructability["design_class"] == "joint")
        & constructability["l_max_kg_per_m"].round(6).eq(
            joint_knees.loc["front_specific_minmax", "l_max_kg_per_m"]
        )
        & constructability["s_max_kg_per_m"].round(6).eq(
            joint_knees.loc["front_specific_minmax", "s_max_kg_per_m"]
        )
    ].iloc[0]
    document.add_paragraph(
        feasibility_sentence
        + " The selected "
        "continuous joint knee had "
        f"{int(discretization_loss['continuous_slope_transitions'])} slope and "
        f"{int(discretization_loss['continuous_surface_transitions'])} surface "
        "transitions and minimum segment length "
        f"{int(continuous_knee_constructability['maximum_minimum_segment_cells'])} cell. "
        "Generic mapping reduced these to "
        f"{int(discretization_loss['mapped_slope_transitions'])} and "
        f"{int(discretization_loss['mapped_surface_transitions'])} transitions; "
        "Lmax changed by "
        f"{discretization_loss['l_max_absolute_change_kg_per_m']:+.2f} kg m-1 "
        f"({discretization_loss['l_max_percent_change']:+.2f}%) and Smax by "
        f"{discretization_loss['s_max_absolute_change_kg_per_m']:+.2f} kg m-1 "
        f"({discretization_loss['s_max_percent_change']:+.2f}%). The mapped profile is "
        "a post hoc rule-compliant translation, not a constrained optimum. Every frozen transition-"
        "penalty setting selected a uniform design. Labor-scarcity and phase-diagram "
        "classifications remain illustrative decision weights, not community or "
        "building-code recommendations."
    )
    jma_path = results / "jma_daily_evaluation.csv"
    if jma_path.exists():
        jma = pd.read_csv(jma_path)
        document.add_heading("3.6 JMA station-winter scenarios", level=2)
        document.add_paragraph(
            f"Daily forcing covered {jma['station_id'].nunique()} stations and "
            f"{len(jma_pairs)} station-winters without missing temperature or snowfall "
            "days. Joint-to-uniform Lmax ratios ranged from "
            f"{jma_summary.loc['joint_to_uniform_l_max_ratio', 'minimum']:.2f} to "
            f"{jma_summary.loc['joint_to_uniform_l_max_ratio', 'maximum']:.2f}, while "
            "Smax reductions ranged from "
            f"{jma_summary.loc['s_max_reduction_percent', 'minimum']:.2f}% to "
            f"{jma_summary.loc['s_max_reduction_percent', 'maximum']:.2f}%. "
            "These paired modeled trade-offs were strongly station-winter dependent. "
            "Table 9 gives the paired station-winter results. "
            "Daily ground observations are supplementary forcing, not roof-scale "
            "validation, and daily Smax is not directly comparable with subdaily "
            "synthetic Smax."
        )

    document.add_heading("4. Discussion", level=1)
    document.add_paragraph(
        "The frozen research question can be answered narrowly: under identical modeled "
        f"forcing, spatial variation retained the {uniform_unique_count} uniform "
        "objective regimes and introduced "
        f"{joint_unique_count} "
        "observed intermediate Lmax-Smax combinations that the exhaustive uniform grid "
        "did not attain under the same objective caps. This is "
        "an expansion of the modeled trade-off set, not universal superiority. The "
        "joint front's larger hypervolume and zero additive epsilon are compatible "
        "because it contains the uniform endpoints while filling intermediate regions."
    )
    document.add_paragraph(
        "Knee selection is unsuitable as the primary cross-front claim: front-specific "
        "normalization and common normalization selected different joint designs. "
        "Matched objective caps provide a clearer interpretation. Robustness also cannot "
        "be inferred from nominal objectives, because nominally equivalent profiles had "
        "widely different upper-tail outcomes and the best scalar score was tied."
    )
    document.add_paragraph(
        "Constructability materially changes the interpretation. The intermediate joint "
        "trade-offs were found in continuous profiles that did not directly satisfy all "
        "segment and transition limits. Post hoc mapping changed the selected knee's "
        f"Lmax by {discretization_loss['l_max_percent_change']:+.2f}% and Smax by "
        f"{discretization_loss['s_max_percent_change']:+.2f}%, moving the point toward "
        "the low-retention, high-release shedding regime; every frozen transition "
        "penalty selected a uniform regime. Thus, the selected continuous-profile "
        "location in objective space was not preserved by the rule-compliant post hoc "
        "mapping. Graded concepts require "
        "constrained reoptimization and physical testing before practical comparison."
    )
    document.add_paragraph(
        "Retention, shedding, active melting, manual removal, and hybrid systems solve "
        "different site problems. Retention can limit routine eave release and reduce "
        "snow-fall-space demand, but requires verified structural capacity, drainage, "
        "waterproofing, inspection, and maintenance. Shedding can reduce retained mass "
        "but requires a controlled receiving area and falling-snow measures. The model "
        "does not invalidate either strategy, and Japanese guidance cannot be generalized "
        "to other climates, building traditions, or regulatory systems without local "
        f"evidence {japanese_citations}."
    )
    document.add_heading("4.1 Limitations", level=2)
    document.add_paragraph(
        "The model has not been calibrated or validated against roof-scale shedding "
        "measurements. It omits wind redistribution, three-dimensional flow, snow "
        "fracture and slab mechanics, cornices, local roof details, heat transfer through "
        "the assembly, impact trajectories, drainage blockage, structural response, and "
        "pedestrian exposure. Friction, adhesion, compaction, melt, and aging parameters "
        f"include explicit assumptions. The {production_step_text}-h production interval "
        "is the finest frozen-grid reference rather than proof of convergence beyond "
        "0.5 h. The 1-h and targeted 0.25-h comparisons showed that event mass, event "
        "count, and SSCI were particularly interval sensitive. The heterogeneous sets are "
        "multi-seed heuristic results rather than guaranteed complete fronts. Daily JMA "
        "observations do not resolve subdaily events, and ground snowfall or snow depth "
        "is not roof snow mass. Accordingly, the study does not establish structural "
        "safety, pedestrian safety, code compliance, injury reduction, lifecycle cost, "
        "or universal effectiveness."
    )
    document.add_heading("4.2 Future validation", level=2)
    document.add_paragraph(
        "Future work should calibrate interface parameters across temperature, liquid-"
        "water content, roughness, aging, and load; validate mass and release timing "
        "against instrumented roofs; add wind redistribution and three-dimensional edge "
        "effects; define an interval-invariant release-rate outcome if cross-timestep "
        "comparison is required; optimize explicitly constrained segments and surface "
        "classes; and predefine external-validation metrics before fitting."
    )

    document.add_heading("5. Conclusions", level=1)
    document.add_paragraph(
        f"An exhaustive uniform reference had {uniform_unique_count} modeled objective "
        "regimes. Heuristic joint spatial grading preserved those regimes and added "
        f"{joint_unique_count} intermediate "
        "trade-off combinations. However, no joint row directly passed all frozen "
        "constructability checks, post hoc mapping shifted the selected point toward the "
        "low-retention, high-release shedding regime, nominal objectives did not "
        "determine robustness, and the targeted 0.25-h check confirmed interval "
        "sensitivity. Constructability and numerical resolution therefore substantially "
        "limit interpretation of the nominal trade-off expansion. "
        "The findings are reproducible design-space hypotheses for constrained "
        "reoptimization and roof-scale validation, not design approval or safety "
        "certification."
    )

    document.add_heading("Declarations", level=1)
    document.add_paragraph(f"Funding: {metadata['funding']}")
    document.add_paragraph(
        f"Declaration of competing interests: {metadata['competing_interests']}"
    )
    document.add_paragraph(
        "Data and code availability: Source code, immutable quantitative inputs and "
        "redistributable public-data snapshots, acquisition metadata, generated result "
        "tables, and reproduction instructions are available at "
        "https://github.com/bougtoir/"
        "functionally-graded-roof-snow-management."
    )
    document.add_heading(
        "Declaration of generative AI and AI-assisted technologies in the manuscript "
        "preparation process",
        level=2,
    )
    document.add_paragraph(
        "During the preparation of this work, the author used Devin "
        "(Cognition AI) to support coding, public-data acquisition, analysis scripting, "
        "and drafting. Before submission, the author must review and edit the content as "
        "needed and replace this sentence with confirmation that the author takes full "
        "responsibility for the content of the published article."
    )
    document.add_paragraph(
        "CRediT authorship contribution statement: REQUIRED—author confirmation of "
        "roles before submission."
    )

    document.add_heading("References", level=1)
    references = references.assign(
        sort_author=references["authors"].str.casefold()
    ).sort_values(["sort_author", "year", "record_id"])
    for _, row in references.iterrows():
        document.add_paragraph(_reference_text(row))

    document.add_page_break()
    document.add_heading("Editable tables", level=1)
    _add_dataframe_table(
        document,
        table_1,
        "Table 1. Audited objective-set summary.",
        [
            "design_class",
            "nondominated_design_rows",
            "unique_objective_pairs",
            "descriptive_knee_l_max_kg_per_m",
            "descriptive_knee_s_max_kg_per_m",
            "normalized_hypervolume_fraction",
            "additive_epsilon_vs_uniform",
        ],
        headers=[
            "Class",
            "Design rows",
            "Objective pairs",
            "Descriptive knee Lmax",
            "Descriptive knee Smax",
            "Norm. HV",
            "Epsilon vs uniform",
        ],
    )
    _add_dataframe_table(
        document,
        table_2,
        "Table 2. Selected joint-profile cell properties.",
        [
            "cell",
            "continuous_slope_deg",
            "discrete_slope_deg",
            "continuous_mu_static",
            "continuous_adhesion_pa",
            "discrete_material_class",
        ],
        headers=[
            "Cell",
            "Continuous slope",
            "Discrete slope",
            "Continuous friction",
            "Continuous adhesion",
            "Generic class",
        ],
    )
    _add_dataframe_table(
        document,
        table_3.assign(
            mapping=table_3["mapping"].replace(
                {"manufacturable_discrete_generic_classes": "discrete_generic"}
            )
        ),
        "Table 3. Post hoc continuous-to-discrete objective-space movement.",
        [
            "mapping",
            "l_max_kg_per_m",
            "s_max_kg_per_m",
            "mean_ssci",
            "mean_event_count",
            "slope_transition_count",
            "surface_transition_count",
            "l_max_change_kg_per_m",
            "l_max_change_percent",
            "s_max_change_kg_per_m",
            "s_max_change_percent",
        ],
        headers=[
            "Mapping",
            "Lmax",
            "Smax",
            "Mean SSCI",
            "Event count",
            "Slope transitions",
            "Surface transitions",
            "Lmax change (kg m-1)",
            "Lmax change (%)",
            "Smax change (kg m-1)",
            "Smax change (%)",
        ],
    )
    document.add_paragraph(
        "Table 4 (strategy phase-diagram classifications) is supplied separately as "
        "an editable CSV because the complete table is too large for useful display "
        "in the manuscript."
    )
    _add_dataframe_table(
        document,
        table_5,
        "Table 5. Frozen labor-scarcity decision choices.",
        [
            "labor_scarcity",
            "selected_design_class",
            "strategy",
            "l_max_kg_per_m",
            "s_max_kg_per_m",
            "mean_manual_triggers",
        ],
        headers=[
            "Labor scarcity",
            "Design class",
            "Strategy",
            "Lmax",
            "Smax",
            "Manual triggers",
        ],
    )
    _add_dataframe_table(
        document,
        pd.concat(
            [
                group.iloc[sorted({0, len(group) // 2, len(group) - 1})]
                for _, group in table_6.groupby("design_class", sort=False)
            ],
            ignore_index=True,
        ),
        "Table 6. Complexity metrics for selected frontier rows.",
        [
            "design_class",
            "l_max_kg_per_m",
            "s_max_kg_per_m",
            "maximum_adjacent_slope_change_deg",
            "transition_count",
            "minimum_segment_cells",
            "meets_minimum_segment_length",
            "meets_maximum_transitions",
        ],
        headers=[
            "Class",
            "Lmax",
            "Smax",
            "Max slope change",
            "Transitions",
            "Min segment cells",
            "Min segment met",
            "Transition limit met",
        ],
    )
    _add_dataframe_table(
        document,
        table_7,
        "Table 7. Complexity-penalty selections.",
        [
            "transition_penalty",
            "design_class",
            "l_max_kg_per_m",
            "s_max_kg_per_m",
            "transition_count",
            "selection_score",
        ],
        headers=[
            "Penalty",
            "Class",
            "Lmax",
            "Smax",
            "Transitions",
            "Score",
        ],
    )
    _add_dataframe_table(
        document,
        table_8,
        "Table 8. Monte Carlo robustness summary for candidate designs.",
        [
            "candidate_id",
            "design_class",
            "nominal_l_max_kg_per_m",
            "nominal_s_max_kg_per_m",
            "quantile_l_max_kg_per_m",
            "quantile_s_max_kg_per_m",
            "probability_manual_intervention",
            "robust_selected",
            "minimum_score_tie",
        ],
        maximum_rows=30,
        headers=[
            "Candidate",
            "Class",
            "Nominal Lmax",
            "Nominal Smax",
            "Q95 Lmax",
            "Q95 Smax",
            "P(manual)",
            "Robust selected",
            "Minimum-score tie",
        ],
    )
    if jma_path.exists():
        table_9 = pd.read_csv(tables / "table_9_jma_daily_evaluation.csv")
        _add_dataframe_table(
            document,
            table_9,
            "Table 9. Supplementary paired JMA station-winter evaluation.",
            [
                "station_id",
                "winter",
                "station_name",
                "uniform_l_max_kg_per_m",
                "joint_l_max_kg_per_m",
                "joint_to_uniform_l_max_ratio",
                "uniform_s_max_kg_per_m",
                "joint_s_max_kg_per_m",
                "s_max_reduction_percent",
            ],
            headers=[
                "Station",
                "Winter",
                "Name",
                "Uniform Lmax",
                "Joint Lmax",
                "Lmax ratio",
                "Uniform Smax",
                "Joint Smax",
                "Smax reduction (%)",
            ],
        )
    document.add_heading("Figure captions", level=1)
    for caption in FIGURE_CAPTIONS:
        document.add_paragraph(caption)

    _enforce_ascii(document)
    output = root / "manuscript" / "build" / "manuscript.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    submission_output = root / "manuscript" / "manuscript_CRST.docx"
    document.save(submission_output)
    document.save(root / "manuscript" / "manuscript_CRST_final.docx")
    document.save(root / "manuscript" / "manuscript_CRST_final_0p5h.docx")
    document.save(
        root / "manuscript" / "manuscript_CRST_submission_final.docx"
    )
    return submission_output


def _find_paragraph(document: Document, text: str):
    for paragraph in document.paragraphs:
        if text in paragraph.text:
            return paragraph
    raise ValueError(f"paragraph not found: {text}")


def _remove_paragraph(paragraph) -> None:
    parent = paragraph._p.getparent()
    if parent is not None:
        parent.remove(paragraph._p)


def _insert_after(paragraph, elements: list) -> None:
    anchor = paragraph._p
    for element in elements:
        anchor.addnext(element)
        anchor = element


def _figure_elements(document: Document, path: Path, caption: str) -> list:
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(12)
    image_paragraph.add_run().add_picture(str(path), width=Inches(6.2))
    caption_paragraph = document.add_paragraph(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(12)
    caption_paragraph.paragraph_format.space_after = Pt(6)
    return [image_paragraph._p, caption_paragraph._p]


def build_inline_manuscript(root: Path) -> Path:
    source = root / "manuscript" / "manuscript_CRST.docx"
    document = Document(source)

    table_4 = pd.read_csv(
        root / "tables" / "generated" / "table_4_strategy_phase_diagram.csv"
    )
    _add_dataframe_table(
        document,
        table_4,
        "Table 4. Strategy phase-diagram classifications "
        "(selected columns; the complete table is supplied as CSV).",
        [
            "snowfall_factor",
            "shedding_penalty_weight",
            "labor_scarcity",
            "allowable_shed_kg_per_m",
            "constraint_feasible",
            "selected_design_class",
            "strategy",
            "mean_manual_triggers",
        ],
        headers=[
            "Snowfall factor",
            "Shed penalty",
            "Labor scarcity",
            "Allowable shed",
            "Feasible",
            "Class",
            "Strategy",
            "Manual triggers",
        ],
    )

    table_elements = {}
    for number in range(1, 10):
        caption = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith(f"Table {number}.")
        )
        table = caption._p.getnext()
        if table is None or table.tag != qn("w:tbl"):
            raise ValueError(f"editable table missing after Table {number} caption")
        caption.paragraph_format.space_before = Pt(12)
        caption.paragraph_format.space_after = Pt(6)
        table_elements[number] = [caption._p, table]

    editable_heading = _find_paragraph(document, "Editable tables")
    preceding = editable_heading._p.getprevious()
    if preceding is not None:
        page_breaks = preceding.findall(f".//{qn('w:br')}")
        if any(item.get(qn("w:type")) == "page" for item in page_breaks):
            preceding.getparent().remove(preceding)
    _remove_paragraph(editable_heading)
    _remove_paragraph(_find_paragraph(document, "Table 4 (strategy phase-diagram"))
    _remove_paragraph(_find_paragraph(document, "Figure captions"))
    for caption in FIGURE_CAPTIONS:
        _remove_paragraph(_find_paragraph(document, caption))

    primary_anchor = _find_paragraph(document, "Figure 1 summarizes")
    convergence_anchor = _find_paragraph(document, "Figure 6 reports")
    robustness_anchor = _find_paragraph(document, "Figures 8-9 and Tables 4-8")
    jma_anchor = _find_paragraph(document, "Table 9 gives")

    figure_elements = {
        number: _figure_elements(
            document,
            root / "figures" / "png" / filename,
            FIGURE_CAPTIONS[number - 1],
        )
        for number, filename in enumerate(FIGURE_FILENAMES, start=1)
    }
    _insert_after(
        primary_anchor,
        [
            *figure_elements[1],
            *figure_elements[2],
            *figure_elements[3],
            *figure_elements[4],
            *figure_elements[5],
            *table_elements[1],
            *table_elements[2],
            *table_elements[3],
        ],
    )
    _insert_after(
        convergence_anchor,
        [*figure_elements[6], *figure_elements[7]],
    )
    _insert_after(
        robustness_anchor,
        [
            *figure_elements[8],
            *figure_elements[9],
            *table_elements[4],
            *table_elements[5],
            *table_elements[6],
            *table_elements[7],
            *table_elements[8],
        ],
    )
    _insert_after(jma_anchor, table_elements[9])

    _enforce_ascii(document)
    output = root / "manuscript" / "manuscript_CRST_inline.docx"
    document.save(output)
    document.save(
        root / "manuscript" / "manuscript_CRST_inline_final.docx"
    )
    document.save(
        root / "manuscript" / "manuscript_CRST_inline_final_0p5h.docx"
    )
    return output


def build_editable_tables(root: Path) -> Path:
    source = Document(root / "manuscript" / "manuscript_CRST_inline.docx")
    document = Document()
    _configure_document(document)
    document.add_heading("Editable manuscript tables", level=0)
    document.add_paragraph(TITLE)
    body = document._body._element
    for number in range(1, 10):
        caption = next(
            paragraph
            for paragraph in source.paragraphs
            if paragraph.text.startswith(f"Table {number}.")
        )
        table = caption._p.getnext()
        if table is None or table.tag != qn("w:tbl"):
            raise ValueError(f"editable table missing after Table {number} caption")
        body.insert(len(body) - 1, deepcopy(caption._p))
        body.insert(len(body) - 1, deepcopy(table))
    _enforce_ascii(document)
    output = root / "manuscript" / "editable_tables_CRST.docx"
    document.save(output)
    return output


def build_editable_figures(root: Path) -> Path:
    presentation = Presentation()
    presentation.slide_width = PptxInches(13.333)
    presentation.slide_height = PptxInches(7.5)
    blank_layout = presentation.slide_layouts[6]

    for number, (filename, caption) in enumerate(
        zip(FIGURE_FILENAMES, FIGURE_CAPTIONS, strict=True),
        start=1,
    ):
        slide = presentation.slides.add_slide(blank_layout)
        title = slide.shapes.add_textbox(
            PptxInches(0.5),
            PptxInches(0.15),
            PptxInches(12.333),
            PptxInches(0.5),
        )
        title_frame = title.text_frame
        title_frame.clear()
        title_paragraph = title_frame.paragraphs[0]
        title_paragraph.text = f"Figure {number}"
        title_paragraph.alignment = PP_ALIGN.CENTER
        title_paragraph.runs[0].font.name = "Arial"
        title_paragraph.runs[0].font.size = PptxPt(24)
        title_paragraph.runs[0].font.bold = True

        image_path = root / "figures" / "png" / filename
        with Image.open(image_path) as image:
            image_width, image_height = image.size
        width = 12.0
        height = width * image_height / image_width
        if height > 5.6:
            height = 5.6
            width = height * image_width / image_height
        slide.shapes.add_picture(
            str(image_path),
            PptxInches((13.333 - width) / 2),
            PptxInches(0.85 + (5.6 - height) / 2),
            width=PptxInches(width),
            height=PptxInches(height),
        )

        caption_box = slide.shapes.add_textbox(
            PptxInches(0.5),
            PptxInches(6.55),
            PptxInches(12.333),
            PptxInches(0.65),
        )
        caption_frame = caption_box.text_frame
        caption_frame.clear()
        caption_paragraph = caption_frame.paragraphs[0]
        caption_paragraph.text = caption
        caption_paragraph.alignment = PP_ALIGN.CENTER
        caption_paragraph.runs[0].font.name = "Arial"
        caption_paragraph.runs[0].font.size = PptxPt(11)

    output = root / "manuscript" / "editable_figures_CRST.pptx"
    presentation.save(output)
    return output


def build_supplement(root: Path) -> Path:
    document = Document()
    _configure_document(document)
    document.add_heading("Supplementary material", level=0)
    document.add_paragraph(TITLE)
    document.add_heading("S1. Frozen analysis and implementation", level=1)
    document.add_paragraph(
        "The analysis plan, machine-readable freeze, and deviation log are included "
        "in the repository. All primary outcomes, optimization seeds, resolution "
        "choices, and classification weights were defined before production results."
    )
    document.add_heading("S2. Reproduction", level=1)
    document.add_paragraph(
        "Create the pinned Python 3.11 environment, run `make lint` and `make test`, "
        "then run `make all`. The production workflow regenerates processed synthetic "
        "and JMA weather, evaluates baselines and numerical resolution, resumes or executes "
        "optimization, runs sensitivity and robustness, creates figures and editable "
        "tables, builds submission documents, validates them, and packages the "
        "submission."
    )
    document.add_heading("S3. Model scope", level=1)
    document.add_paragraph(
        "Cell mass balance is snowfall input minus melt and eave release. Static "
        "friction and adhesion define onset; kinetic friction defines transport. "
        "The model is deterministic for a fixed design and weather series. It is a "
        "threshold-release model rather than a structural, fracture-mechanics, or "
        "computational-fluid-dynamics model."
    )
    document.add_heading("S4. Public-data provenance", level=1)
    ledger = pd.read_csv(root / "data" / "metadata" / "acquisition_ledger.csv")
    counts = ledger["local_status"].value_counts().to_dict()
    document.add_paragraph(
        "The consolidated ledger contains "
        f"{len(ledger)} records: "
        + ", ".join(f"{key}={value}" for key, value in counts.items())
        + ". Each locally retained record includes a path, size, checksum, retrieval "
        "time, request context where applicable, and usage-conditions reference. "
        "Records marked not_recovered identify non-quantitative source leads that were "
        "not retained locally and were not used as analysis inputs. Quantitative inputs "
        "and redistributable evidence are retained in verified repository snapshots; "
        "non-redistributed third-party documents remain identified by provenance "
        "metadata."
    )
    document.add_heading("S5. Machine-readable outputs", level=1)
    document.add_paragraph(
        "Complete design-space, Pareto, per-seed history, sensitivity, robustness, "
        "JMA scenario, phase-diagram, and material-mapping tables are supplied as CSV. "
        "The selected manuscript values are generated directly from these files."
    )
    _enforce_ascii(document)
    output = root / "manuscript" / "build" / "supplement.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    submission_output = root / "manuscript" / "supplement_CRST.docx"
    document.save(submission_output)
    document.save(root / "manuscript" / "supplement_CRST_final.docx")
    document.save(root / "manuscript" / "supplementary_material_CRST.docx")
    document.save(
        root
        / "manuscript"
        / "supplementary_material_CRST_submission_final.docx"
    )
    return submission_output


def build_cover_letter(root: Path) -> Path:
    metadata = yaml.safe_load(
        (root / "manuscript" / "author_metadata.yaml").read_text(encoding="utf-8")
    )
    author = next(item for item in metadata["authors"] if item.get("corresponding"))
    document = Document()
    _configure_document(document)
    document.add_paragraph("Editor-in-Chief")
    document.add_paragraph("Cold Regions Science and Technology")
    document.add_paragraph("Dear Editor,")
    document.add_paragraph(
        f"Please consider my manuscript, “{TITLE},” as a Research Article in "
        "Cold Regions Science and Technology."
    )
    document.add_paragraph(
        "The manuscript presents a reproducible reduced-order comparison of optimized "
        "spatially uniform and functionally graded passive roof-snow strategies. It "
        "maps the retained-mass versus discrete-release trade-off, includes geometry and "
        "surface ablations, numerical-resolution checks, sensitivity, Monte Carlo "
        "perturbation, manufacturing-complexity analysis, and supplementary scenarios "
        "from official Japanese weather observations. Spatial grading produced four "
        "nominal intermediate trade-offs absent from the exhaustive uniform grid, but "
        "constructability and timestep sensitivity substantially qualify their "
        "interpretation."
    )
    document.add_paragraph(
        "The work fits the journal because it addresses cold-regions roof-snow "
        "processes, explicitly positions the model against roof-snow literature, and "
        "states the need for empirical validation. It makes no claim of structural "
        "safety, code compliance, pedestrian safety, or universal performance."
    )
    document.add_paragraph(
        "REQUIRED before submission: corresponding-author confirmation that the "
        "manuscript is original, is not under consideration elsewhere, and has been "
        "approved by all authors."
    )
    document.add_paragraph(
        "All source code, quantitative inputs, redistributable public-data snapshots, "
        "checksum ledgers, and generated analysis artifacts are supplied in a public "
        "reproducibility repository."
    )
    document.add_paragraph("Sincerely,")
    document.add_paragraph(author["name"])
    document.add_paragraph(author["affiliation"])
    document.add_paragraph(author["postal_address"])
    document.add_paragraph(author["email"])
    _enforce_ascii(document)
    output = root / "manuscript" / "build" / "cover_letter.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    submission_output = root / "manuscript" / "cover_letter_CRST.docx"
    document.save(submission_output)
    document.save(root / "manuscript" / "cover_letter_CRST_final.docx")
    document.save(
        root / "manuscript" / "cover_letter_CRST_submission_final.docx"
    )
    return submission_output


def build_text_files(root: Path) -> list[Path]:
    build = root / "manuscript" / "build"
    build.mkdir(parents=True, exist_ok=True)
    reproduction_summary_path = (
        root / "audit" / "0p5h_revision" / "reproduction_summary.json"
    )
    if reproduction_summary_path.exists():
        reproduction = json.loads(
            reproduction_summary_path.read_text(encoding="utf-8")
        )
        reproduction_text = (
            "A detached clean-environment run regenerated "
            f"{reproduction['identical_files']} of "
            f"{reproduction['compared_files']} quantitative CSV files byte for byte; "
            f"Ruff and {reproduction['test_count']} tests passed. "
        )
    else:
        reproduction_text = (
            "The current 0.5-h checkpoint manifests passed configuration, source, "
            "weather, seed, and checksum compatibility checks. "
        )
    files: dict[str, str] = {
        "highlights.txt": (
            "- Grading added four nominal trade-offs absent from the uniform grid.\n"
            "- No joint row passed all frozen constructability checks.\n"
            "- Post hoc mapping shifted the selected point toward shedding.\n"
            "- Timestep sensitivity limits interpretation beyond the 0.5-h reference.\n"
        ),
        "scope_fit.md": (
            "# Scope fit\n\n"
            "The study addresses roof-snow accumulation and passive release in cold "
            "regions using reproducible computational modeling, official weather "
            "observations, and explicit links to roof-snow deposition, interface, and "
            "validation literature published in *Cold Regions Science and Technology*. "
            "The submission is framed as reduced-order design-space analysis and does "
            "not claim structural validation.\n"
        ),
        "crst_checklist.md": (
            "# CRST submission checklist\n\n"
            "- [x] Editable single-column Word manuscript\n"
            "- [x] Abstract no longer than 250 words\n"
            "- [x] Six keywords\n"
            "- [x] Separate highlights file with 3-5 bullets, each <=85 characters\n"
            "- [x] Continuous line numbering requested in DOCX XML\n"
            "- [x] Editable manuscript tables\n"
            "- [x] Figures supplied separately as EPS, PDF, SVG, TIFF, and PNG\n"
            "- [x] Data/code availability and generative-AI statements\n"
            "- [ ] Author affiliation and postal address confirmed\n"
            "- [ ] Funding and competing-interest declarations confirmed\n"
            "- [ ] CRediT roles confirmed by the author\n"
            "- [ ] Originality and author-approval statement confirmed\n"
            "- [x] Current official Guide for Authors rechecked and archived\n"
        ),
        "reproducibility_readme.md": (
            "# Reproducibility\n\n"
            "Use Python 3.11 and install the pinned project plus development "
            "dependencies with `python -m pip install -e '.[dev]'`. Run `make lint`, "
            "`make test`, and `make all` from the repository root. The production "
            "pipeline regenerates quantitative inputs derived from retained raw "
            "snapshots, optimization outputs, figures, tables, manuscript files, "
            "validation reports, and this submission package.\n\n"
            + reproduction_text
            + "The five public-source snapshots used by the final literature audit "
            "are retained under `data/raw/published_sources/` with ledgered URLs, "
            "sizes, SHA-256 values, and usage conditions. They are not quantitative "
            "analysis inputs. The targeted finalization intentionally did not repeat "
            "a full clean pipeline rebuild after the one-time Figure and Table "
            "regeneration; see `FINAL_HANDOFF.md` and `manuscript_values.csv`.\n"
        ),
    }
    outputs = []
    for name, content in files.items():
        path = build / name
        path.write_text(content, encoding="utf-8")
        outputs.append(path)
    aliases = {
        "highlights.txt": "highlights_CRST.txt",
        "scope_fit.md": "CRST_scope_fit.md",
        "crst_checklist.md": "CRST_submission_checklist.md",
        "reproducibility_readme.md": "REPRODUCIBILITY_README.md",
    }
    for source_name, alias_name in aliases.items():
        path = root / "manuscript" / alias_name
        path.write_text(files[source_name], encoding="utf-8")
        outputs.append(path)
    final_aliases = {
        "scope_fit.md": "CRST_scope_fit_final.md",
        "crst_checklist.md": "CRST_submission_checklist_final.md",
    }
    for source_name, alias_name in final_aliases.items():
        path = root / "manuscript" / alias_name
        path.write_text(files[source_name], encoding="utf-8")
        outputs.append(path)
    return outputs


def package_submission(root: Path) -> Path:
    output = root / "submission" / "CRST_submission_package_FINAL.zip"
    output.parent.mkdir(parents=True, exist_ok=True)
    submission_files = [
        (
            root / "manuscript" / "manuscript_CRST_submission_final.docx",
            "manuscript_CRST_submission_final.docx",
        ),
        (
            root / "manuscript" / "manuscript_CRST_inline_final_0p5h.docx",
            "review_copy/manuscript_CRST_inline_final_0p5h.docx",
        ),
        (
            root / "manuscript" / "cover_letter_CRST_submission_final.docx",
            "cover_letter_CRST_submission_final.docx",
        ),
        (
            root
            / "manuscript"
            / "supplementary_material_CRST_submission_final.docx",
            "supplementary_material_CRST_submission_final.docx",
        ),
        (
            root / "manuscript" / "editable_tables_CRST.docx",
            "editable_tables_CRST.docx",
        ),
        (
            root / "manuscript" / "editable_figures_CRST.pptx",
            "editable_figures_CRST.pptx",
        ),
        (
            root / "manuscript" / "highlights_CRST.txt",
            "highlights_CRST.txt",
        ),
        (
            root / "manuscript" / "CRST_submission_checklist_final.md",
            "CRST_submission_checklist_final.md",
        ),
        (
            root / "manuscript" / "CRST_scope_fit_final.md",
            "CRST_scope_fit_final.md",
        ),
        (
            root / "manuscript" / "REPRODUCIBILITY_README.md",
            "REPRODUCIBILITY_README.md",
        ),
        (
            root / "references" / "reference_audit.csv",
            "audit/reference_audit.csv",
        ),
        (
            root / "audit" / "FINAL_AUDIT.md",
            "audit/FINAL_AUDIT.md",
        ),
        (
            root / "audit" / "FINAL_FINISHING_AUDIT.md",
            "audit/FINAL_FINISHING_AUDIT.md",
        ),
        (
            root / "audit" / "REPRODUCIBILITY_AUDIT.md",
            "audit/REPRODUCIBILITY_AUDIT.md",
        ),
        (
            root / "audit" / "FABRICATION_AUDIT.md",
            "audit/FABRICATION_AUDIT.md",
        ),
        (
            root / "audit" / "FABRICATION_AUDIT_FINAL.md",
            "audit/FABRICATION_AUDIT_FINAL.md",
        ),
        (
            root / "audit" / "NUMERICAL_CONSISTENCY_FINAL.md",
            "audit/NUMERICAL_CONSISTENCY_FINAL.md",
        ),
        (
            root / "references" / "REFERENCE_AUDIT_FINAL.csv",
            "audit/REFERENCE_AUDIT_FINAL.csv",
        ),
        (
            root / "audit" / "FINAL_CRST_REVIEW.md",
            "audit/FINAL_CRST_REVIEW.md",
        ),
        (
            root / "audit" / "0p5h_revision" / "TIMESTEP_AUDIT.md",
            "audit/TIMESTEP_AUDIT.md",
        ),
        (
            root / "audit" / "0p5h_revision" / "NUMERICAL_REFERENCE_AUDIT.md",
            "audit/NUMERICAL_REFERENCE_AUDIT.md",
        ),
        (
            root / "audit" / "0p5h_revision" / "OLD_VS_NEW_PRIMARY_AUDIT.md",
            "audit/OLD_VS_NEW_PRIMARY_AUDIT.md",
        ),
        (
            root / "audit" / "CRST_FORMAT_LANGUAGE_AUDIT.md",
            "audit/CRST_FORMAT_LANGUAGE_AUDIT.md",
        ),
        (
            root / "audit" / "final_revision" / "FINAL_HOSTILE_REVIEW.md",
            "audit/FINAL_HOSTILE_REVIEW.md",
        ),
        (root / "README.md", "repository_README.md"),
        (root / "analysis_freeze.yaml", "analysis_freeze.yaml"),
        (root / "PROJECT_STATE.json", "PROJECT_STATE.json"),
    ]
    include_roots = [
        root / "figures" / "png",
        root / "figures" / "tiff",
        root / "figures" / "vector",
        root / "tables" / "generated",
    ]
    entries = list(submission_files)
    fresh_comparison = (
        root / "audit" / "0p5h_revision" / "fresh_reproduction_comparison.csv"
    )
    if fresh_comparison.exists():
        entries.append(
            (fresh_comparison, "audit/fresh_reproduction_comparison.csv")
        )
    for include_root in include_roots:
        for path in sorted(include_root.rglob("*")):
            if path.is_file():
                entries.append((path, str(path.relative_to(root))))
    manifest_path = root / "submission" / "CRST_submission_manifest.csv"
    with manifest_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["archive_path", "size_bytes", "sha256"],
            lineterminator="\n",
        )
        writer.writeheader()
        for path, archive_path in entries:
            writer.writerow(
                {
                    "archive_path": archive_path,
                    "size_bytes": path.stat().st_size,
                    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                }
            )
    entries.append((manifest_path, manifest_path.name))
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path, archive_path in entries:
            archive.write(path, archive_path)
    legacy_output = root / "submission" / "CRST_submission_package.zip"
    legacy_output.write_bytes(output.read_bytes())
    final_legacy_output = (
        root / "submission" / "CRST_submission_package_final.zip"
    )
    final_legacy_output.write_bytes(output.read_bytes())
    return output


def write_manifest(root: Path, paths: list[Path]) -> Path:
    manifest = []
    for path in paths:
        manifest.append(
            {
                "path": str(path.relative_to(root)),
                "size_bytes": path.stat().st_size,
            }
        )
    output = root / "manuscript" / "build" / "artifact_manifest.json"
    output.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return output


def write_csv_rows(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=list(rows[0]),
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(rows)
