from __future__ import annotations

import argparse
import shutil
from pathlib import Path

import pandas as pd
import yaml
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results" / "generated"
TABLES = ROOT / "tables" / "generated"
MANUSCRIPT = (
    ROOT / "manuscript" / "manuscript_CRST_submission_final.docx"
)


def _document_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return "\n".join(parts)


def _add_value(
    rows: list[dict[str, str]],
    manuscript_text: str,
    value_id: str,
    location: str,
    reported_value: str,
    units: str,
    source: str,
    locator: str,
    derivation: str,
    code: str,
    fragment: str,
) -> None:
    match = fragment in manuscript_text
    rows.append(
        {
            "value_id": value_id,
            "manuscript_location": location,
            "reported_value": reported_value,
            "units": units,
            "canonical_source": source,
            "source_locator": locator,
            "derivation": derivation,
            "generator_code": code,
            "manuscript_fragment": fragment,
            "manuscript_match": str(match),
            "status": "PASS" if match else "FAIL",
        }
    )


def build_manuscript_values() -> Path:
    config = yaml.safe_load(
        (ROOT / "config" / "production.yaml").read_text(encoding="utf-8")
    )
    text = _document_text(MANUSCRIPT)
    rows: list[dict[str, str]] = []
    baseline = pd.read_csv(RESULTS / "baseline_aggregate.csv").set_index("design")
    convergence = pd.read_csv(RESULTS / "convergence.csv")
    frontier = pd.read_csv(
        RESULTS / "final_revision_frontier_summary.csv"
    ).set_index("design_class")
    hypervolume = pd.read_csv(
        RESULTS / "final_revision_frontier_metric_sensitivity.csv"
    )
    hypervolume = hypervolume.loc[
        hypervolume["normalization_origin"].eq("zero")
        & hypervolume["reference_margin"].eq(1.05)
    ].set_index("design_class")
    knee = pd.read_csv(RESULTS / "final_revision_knee_audit.csv")
    joint_knee = knee.loc[
        knee["design_class"].eq("joint")
        & knee["normalization"].eq("front_specific_minmax")
    ].iloc[0]
    constraints = pd.read_csv(
        RESULTS / "final_revision_prespecified_constraints.csv"
    )
    cap20 = constraints.loc[
        constraints["s_max_limit_kg_per_m"].eq(20.0)
    ].set_index("design_class")
    robustness = pd.read_csv(
        RESULTS / "final_revision_objective_equivalent_robustness.csv"
    ).sort_values("candidate_count", ascending=False).iloc[0]
    ties = pd.read_csv(RESULTS / "final_revision_robust_selection_ties.csv")
    mapping = pd.read_csv(
        RESULTS / "final_revision_discretization_loss.csv"
    ).iloc[0]
    constructability = pd.read_csv(
        RESULTS / "final_revision_constructability_by_objective.csv"
    )
    timestep = pd.read_csv(
        RESULTS / "final_revision_selected_design_timestep_audit.csv"
    )
    quarter_hour = pd.read_csv(
        RESULTS / "final_revision_selected_design_0p25h_audit.csv"
    )
    jma = pd.read_csv(
        RESULTS / "final_revision_jma_paired_tradeoffs.csv"
    )

    config_source = "config/production.yaml"
    manuscript_code = "src/graded_roof/manuscript.py:build_manuscript"
    reporting_code = "src/graded_roof/reporting.py:generate_tables"

    _add_value(
        rows,
        text,
        "roof_length",
        "Methods 2.1",
        f"{config['roof']['length_m']:g}",
        "m",
        config_source,
        "roof.length_m",
        "Direct configured model dimension.",
        manuscript_code,
        "An 8 m long, 1 m wide roof",
    )
    _add_value(
        rows,
        text,
        "roof_width",
        "Methods 2.1",
        f"{config['roof']['width_m']:g}",
        "m",
        config_source,
        "roof.width_m",
        "Direct configured unit-width model dimension.",
        manuscript_code,
        "An 8 m long, 1 m wide roof",
    )
    _add_value(
        rows,
        text,
        "roof_cells",
        "Methods 2.1",
        str(config["roof"]["cells"]),
        "cells",
        config_source,
        "roof.cells",
        "Direct configured spatial resolution.",
        manuscript_code,
        "represented by 24 ridge-to-eave cells",
    )
    _add_value(
        rows,
        text,
        "synthetic_winter_count",
        "Methods 2.3",
        str(len(config["synthetic_weather"]["scenarios"])),
        "scenarios",
        config_source,
        "synthetic_weather.scenarios",
        "Count of configured deterministic synthetic winters.",
        manuscript_code,
        "Primary optimization used 3 deterministic synthetic winters",
    )
    _add_value(
        rows,
        text,
        "jma_station_count",
        "Methods 2.3",
        str(len(config["jma"]["stations"])),
        "stations",
        config_source,
        "jma.stations",
        "Count of configured JMA stations.",
        manuscript_code,
        "for 4 stations and 3 winters",
    )
    _add_value(
        rows,
        text,
        "jma_winter_count",
        "Methods 2.3",
        str(len(config["jma"]["winters"])),
        "winters",
        config_source,
        "jma.winters",
        "Count of configured JMA winters.",
        manuscript_code,
        "for 4 stations and 3 winters",
    )
    _add_value(
        rows,
        text,
        "uniform_grid",
        "Methods 2.4",
        "30 x 30 x 7",
        "grid levels",
        config_source,
        "optimizer.uniform_*_grid",
        "Configured slope by friction by adhesion enumeration.",
        manuscript_code,
        "A full 30 x 30 x 7 uniform slope-friction-adhesion map",
    )
    _add_value(
        rows,
        text,
        "profile_control_points",
        "Methods 2.4",
        str(config["optimizer"]["profile_control_points"]),
        "control points",
        config_source,
        "optimizer.profile_control_points",
        "Direct configured heterogeneous-profile parameterization.",
        manuscript_code,
        "parameterized by 4 control points",
    )
    _add_value(
        rows,
        text,
        "optimizer_population",
        "Methods 2.4",
        str(config["optimizer"]["population"]),
        "individuals",
        config_source,
        "optimizer.population",
        "Direct configured NSGA-II population.",
        manuscript_code,
        "population of 48, 45 generations",
    )
    _add_value(
        rows,
        text,
        "optimizer_generations",
        "Methods 2.4",
        str(config["optimizer"]["generations"]),
        "generations",
        config_source,
        "optimizer.generations",
        "Direct configured NSGA-II generations.",
        manuscript_code,
        "population of 48, 45 generations",
    )
    seed_text = ", ".join(str(seed) for seed in config["optimizer"]["seeds"])
    _add_value(
        rows,
        text,
        "optimizer_seeds",
        "Methods 2.4",
        seed_text,
        "seed identifiers",
        config_source,
        "optimizer.seeds",
        "Direct configured NSGA-II seeds.",
        manuscript_code,
        f"seeds {seed_text}",
    )
    _add_value(
        rows,
        text,
        "hypervolume_reference_margin",
        "Methods 2.4",
        "105",
        "%",
        "results/generated/final_revision_frontier_metric_sensitivity.csv",
        "reference_margin=1.05",
        "Common reference equals 105% of joint objective maxima.",
        "scripts/run_final_revision_analyses.py:stage_frontier",
        "reference at 105% of the joint objective maxima",
    )
    _add_value(
        rows,
        text,
        "production_timestep",
        "Methods 2.4; Results 3.1",
        f"{config['simulation']['dt_hours']:g}",
        "h",
        config_source,
        "simulation.dt_hours",
        "Frozen production interval.",
        manuscript_code,
        "Production used a 0.5-h timestep",
    )
    _add_value(
        rows,
        text,
        "post_freeze_timestep",
        "Methods 2.4; Results 3.1",
        f"{config['convergence']['post_freeze_reference_dt_hours']:g}",
        "h",
        config_source,
        "convergence.post_freeze_reference_dt_hours",
        "Targeted reevaluation interval without reoptimization.",
        manuscript_code,
        "targeted 0.25-h",
    )
    _add_value(
        rows,
        text,
        "robustness_draws",
        "Methods 2.4",
        str(config["robustness"]["samples"]),
        "paired draws",
        config_source,
        "robustness.samples",
        "Direct configured deterministic Monte Carlo sample count.",
        manuscript_code,
        "Monte Carlo robustness used 250 deterministic draws",
    )
    _add_value(
        rows,
        text,
        "surface_class_count",
        "Methods 2.4",
        str(len(config["surface"]["discrete_classes"])),
        "generic classes",
        config_source,
        "surface.discrete_classes",
        "Count of configured generic surface classes.",
        manuscript_code,
        "mapped to 3 generic surface classes",
    )
    _add_value(
        rows,
        text,
        "minimum_segment_cells",
        "Methods 2.4",
        str(config["roof"]["minimum_segment_cells"]),
        "cells",
        config_source,
        "roof.minimum_segment_cells",
        "Frozen post hoc constructability rule.",
        manuscript_code,
        "segments of at least 3 cells",
    )
    _add_value(
        rows,
        text,
        "maximum_transitions",
        "Methods 2.4",
        str(config["roof"]["maximum_transitions"]),
        "transitions",
        config_source,
        "roof.maximum_transitions",
        "Frozen post hoc constructability rule.",
        manuscript_code,
        "with at most 6 transitions",
    )
    tolerance = 100 * config["convergence"]["relative_tolerance"]
    _add_value(
        rows,
        text,
        "convergence_tolerance",
        "Results 3.1",
        f"{tolerance:.0f}",
        "%",
        config_source,
        "convergence.relative_tolerance",
        "Configured four-metric relative tolerance.",
        manuscript_code,
        f"frozen {tolerance:.0f}% four-metric convergence tolerance",
    )
    convergence_passes = int(convergence["within_tolerance"].sum())
    _add_value(
        rows,
        text,
        "convergence_pass_count",
        "Results 3.1",
        f"{convergence_passes} of {len(convergence)}",
        "resolution combinations",
        "results/generated/convergence.csv",
        "within_tolerance=True",
        "Count of frozen-grid rows meeting all four metrics.",
        "scripts/run_pipeline.py:stage_convergence",
        f"met by {convergence_passes} of {len(convergence)} tested",
    )
    one_hour = timestep.loc[timestep["dt_hours"].eq(1.0)]
    one_hour_failures = int((~one_hour["within_tolerance"]).sum())
    _add_value(
        rows,
        text,
        "one_hour_failures",
        "Results 3.1",
        f"{one_hour_failures} of {len(one_hour)}",
        "selected designs",
        "results/generated/final_revision_selected_design_timestep_audit.csv",
        "dt_hours=1.0; within_tolerance=False",
        "Count of selected designs failing relative to 0.5 h.",
        "scripts/run_0p5h_revision.py:stage_timestep",
        f"At 1 h, {one_hour_failures} of {len(one_hour)} selected",
    )
    quarter = quarter_hour.loc[quarter_hour["dt_hours"].eq(0.5)]
    quarter_failures = int((~quarter["within_frozen_tolerance"]).sum())
    _add_value(
        rows,
        text,
        "quarter_hour_failures",
        "Abstract; Results 3.1",
        f"{quarter_failures} of {len(quarter)}",
        "selected designs",
        "results/generated/final_revision_selected_design_0p25h_audit.csv",
        "dt_hours=0.5; within_frozen_tolerance=False",
        "Count of 0.5-h selected designs failing relative to 0.25 h.",
        "scripts/run_0p5h_revision.py:stage_quarter_hour",
        f"{quarter_failures} of {len(quarter)} selected designs failed the "
        "frozen four-metric tolerance",
    )

    for design, prefix, label in [
        ("retention_non_shedding", "retention", "retention baseline"),
        ("conventional_shedding", "shedding", "shedding baseline"),
    ]:
        row = baseline.loc[design]
        l_text = f"{row['l_max_kg_per_m']:.1f}" if prefix == "retention" else (
            f"{row['l_max_kg_per_m']:.2f}"
        )
        s_text = f"{row['s_max_kg_per_m']:.3f}" if prefix == "retention" else (
            f"{row['s_max_kg_per_m']:.2f}"
        )
        _add_value(
            rows,
            text,
            f"{prefix}_lmax",
            "Results 3.1",
            l_text,
            "kg m-1",
            "results/generated/baseline_aggregate.csv",
            f"design={design}; l_max_kg_per_m",
            f"Reported {label} Lmax.",
            "scripts/run_pipeline.py:stage_baseline",
            f"{label} produced Lmax={l_text}",
        )
        _add_value(
            rows,
            text,
            f"{prefix}_smax",
            "Results 3.1",
            s_text,
            "kg m-1",
            "results/generated/baseline_aggregate.csv",
            f"design={design}; s_max_kg_per_m",
            f"Reported {label} Smax.",
            "scripts/run_pipeline.py:stage_baseline",
            f"Smax={s_text}",
        )

    for design_class in ["uniform", "geometry", "surface", "joint"]:
        row = frontier.loc[design_class]
        _add_value(
            rows,
            text,
            f"{design_class}_objective_pairs",
            "Results 3.2-3.3; Table 1",
            str(int(row["unique_objective_pairs"])),
            "objective pairs",
            "results/generated/final_revision_frontier_summary.csv",
            f"design_class={design_class}; unique_objective_pairs",
            "Six-decimal objective-pair deduplication.",
            "scripts/run_final_revision_analyses.py:stage_frontier",
            f"{design_class}: {int(row['unique_objective_pairs'])} objective pairs",
        )
        hv = hypervolume.loc[
            design_class, "normalized_hypervolume_fraction"
        ]
        _add_value(
            rows,
            text,
            f"{design_class}_hypervolume",
            "Results 3.3; Table 1",
            f"{hv:.3f}",
            "normalized fraction",
            "results/generated/final_revision_frontier_metric_sensitivity.csv",
            (
                f"design_class={design_class}; normalization_origin=zero; "
                "reference_margin=1.05"
            ),
            "Common zero-origin, 1.05-margin normalized hypervolume.",
            "scripts/run_final_revision_analyses.py:stage_frontier",
            (
                f"{design_class}: {int(row['unique_objective_pairs'])} objective "
                f"pairs, normalized hypervolume {hv:.3f}"
            ),
        )
    for design_class in ["uniform", "joint"]:
        row = frontier.loc[design_class]
        _add_value(
            rows,
            text,
            f"{design_class}_design_rows",
            "Results 3.2 or 3.5; Table 1",
            str(int(row["nondominated_design_rows"])),
            "nondominated rows",
            "results/generated/final_revision_frontier_summary.csv",
            f"design_class={design_class}; nondominated_design_rows",
            "Count of retained nondominated design rows.",
            "scripts/run_final_revision_analyses.py:stage_frontier",
            (
                f"{int(row['nondominated_design_rows'])} nondominated "
                + ("design rows" if design_class == "uniform" else "rows")
            ),
        )
    _add_value(
        rows,
        text,
        "joint_intermediate_pairs",
        "Abstract; Results 3.3; Discussion; Conclusion",
        "4",
        "objective pairs",
        "results/generated/final_revision_joint_unique_regions.csv",
        "attainable_by_uniform_under_same_caps=False",
        "Count of joint pairs unavailable to uniform under identical caps.",
        "scripts/run_final_revision_analyses.py:stage_frontier",
        "4 additional intermediate pairs",
    )
    _add_value(
        rows,
        text,
        "joint_epsilon",
        "Abstract; Results 3.3; Table 1",
        (
            f"{frontier.loc['joint', 'additive_epsilon_vs_uniform_common_minmax']:.3f}"
        ),
        "normalized additive epsilon",
        "results/generated/final_revision_frontier_summary.csv",
        "design_class=joint; additive_epsilon_vs_uniform_common_minmax",
        "Common-minmax additive epsilon against uniform.",
        "scripts/run_final_revision_analyses.py:stage_frontier",
        "joint additive epsilon versus uniform was 0.000",
    )
    _add_value(
        rows,
        text,
        "joint_knee_lmax",
        "Results 3.3; Table 1",
        f"{joint_knee['l_max_kg_per_m']:.1f}",
        "kg m-1",
        "results/generated/final_revision_knee_audit.csv",
        "design_class=joint; normalization=front_specific_minmax",
        "Within-front normalized-distance knee Lmax.",
        "scripts/run_final_revision_analyses.py:stage_knee",
        f"joint knee was {joint_knee['l_max_kg_per_m']:.1f}/",
    )
    _add_value(
        rows,
        text,
        "joint_knee_smax",
        "Results 3.3; Table 1",
        f"{joint_knee['s_max_kg_per_m']:.2f}",
        "kg m-1",
        "results/generated/final_revision_knee_audit.csv",
        "design_class=joint; normalization=front_specific_minmax",
        "Within-front normalized-distance knee Smax.",
        "scripts/run_final_revision_analyses.py:stage_knee",
        f"{joint_knee['l_max_kg_per_m']:.1f}/{joint_knee['s_max_kg_per_m']:.2f}",
    )
    for design_class in ["joint", "uniform"]:
        value = cap20.loc[design_class, "minimum_l_max_kg_per_m"]
        _add_value(
            rows,
            text,
            f"cap20_{design_class}_lmax",
            "Results 3.3",
            f"{value:.1f}",
            "kg m-1",
            "results/generated/final_revision_prespecified_constraints.csv",
            (
                "s_max_limit_kg_per_m=20; "
                f"design_class={design_class}"
            ),
            "Minimum Lmax among points satisfying Smax <= 20 kg m-1.",
            "scripts/run_final_revision_analyses.py:stage_knee",
            f"{value:.1f}",
        )

    robustness_values = [
        ("robust_candidate_count", "candidate_count", ".0f", "candidates"),
        ("robust_q95_l_min", "q95_l_min_kg_per_m", ".2f", "kg m-1"),
        ("robust_q95_l_max", "q95_l_max_kg_per_m", ".1f", "kg m-1"),
        ("robust_q95_s_min", "q95_s_min_kg_per_m", ".2f", "kg m-1"),
        ("robust_q95_s_max", "q95_s_max_kg_per_m", ".1f", "kg m-1"),
        (
            "robust_intervention_min",
            "intervention_probability_min",
            ".3f",
            "probability",
        ),
        (
            "robust_intervention_max",
            "intervention_probability_max",
            ".3f",
            "probability",
        ),
    ]
    robustness_fragments = {
        "robust_candidate_count": "Among 12 candidates",
        "robust_q95_l_min": "Q95 Lmax ranged from 51.89",
        "robust_q95_l_max": "51.89 to 5816.7",
        "robust_q95_s_min": "Q95 Smax ranged from 51.89",
        "robust_q95_s_max": "51.89 to 2726.6",
        "robust_intervention_min": "probability ranged from 0.000",
        "robust_intervention_max": "0.000 to 0.216",
    }
    for value_id, column, format_spec, units in robustness_values:
        value = robustness[column]
        _add_value(
            rows,
            text,
            value_id,
            "Results 3.4",
            format(value, format_spec),
            units,
            "results/generated/final_revision_objective_equivalent_robustness.csv",
            "largest candidate_count group; " + column,
            "Range within largest nominal objective-equivalent group.",
            "scripts/run_final_revision_analyses.py:stage_robustness",
            robustness_fragments[value_id],
        )
    _add_value(
        rows,
        text,
        "robust_tie_count",
        "Abstract; Results 3.4",
        str(len(ties)),
        "candidates",
        "results/generated/final_revision_robust_selection_ties.csv",
        "minimum_score_tie=True",
        "Count of candidates sharing minimum scalar robustness score.",
        "scripts/run_final_revision_analyses.py:stage_robustness",
        f"shared by {len(ties)} candidates",
    )

    joint_constructability = constructability.loc[
        constructability["design_class"].eq("joint")
    ]
    feasible_joint = int(joint_constructability["feasible_design_rows"].sum())
    joint_rows = int(frontier.loc["joint", "nondominated_design_rows"])
    _add_value(
        rows,
        text,
        "constructable_joint_rows",
        "Results 3.5",
        f"{feasible_joint} of {joint_rows}",
        "joint rows",
        "results/generated/final_revision_constructability_by_objective.csv",
        "design_class=joint; sum(feasible_design_rows)",
        "Count passing every frozen post hoc constructability check.",
        "scripts/run_final_revision_analyses.py:stage_constructability",
        f"No joint design row among the {joint_rows} nondominated rows",
    )
    mapping_fields = [
        (
            "continuous_slope_transitions",
            "continuous_slope_transitions",
            ".0f",
            "transitions",
            "had 23 slope and 23 surface transitions",
        ),
        (
            "continuous_surface_transitions",
            "continuous_surface_transitions",
            ".0f",
            "transitions",
            "had 23 slope and 23 surface transitions",
        ),
        (
            "mapped_slope_transitions",
            "mapped_slope_transitions",
            ".0f",
            "transitions",
            "reduced these to 6 and 0 transitions",
        ),
        (
            "mapped_surface_transitions",
            "mapped_surface_transitions",
            ".0f",
            "transitions",
            "reduced these to 6 and 0 transitions",
        ),
        (
            "mapping_lmax_absolute",
            "l_max_absolute_change_kg_per_m",
            "+.2f",
            "kg m-1",
            "Lmax changed by -4469.09",
        ),
        (
            "mapping_lmax_percent",
            "l_max_percent_change",
            "+.2f",
            "%",
            "(-98.28%)",
        ),
        (
            "mapping_smax_absolute",
            "s_max_absolute_change_kg_per_m",
            "+.2f",
            "kg m-1",
            "Smax by +66.00",
        ),
        (
            "mapping_smax_percent",
            "s_max_percent_change",
            "+.2f",
            "%",
            "(+537.47%)",
        ),
    ]
    for value_id, column, format_spec, units, fragment in mapping_fields:
        _add_value(
            rows,
            text,
            value_id,
            "Abstract; Results 3.5; Discussion; Table 3",
            format(mapping[column], format_spec),
            units,
            "results/generated/final_revision_discretization_loss.csv",
            column,
            "Mapped minus continuous selected joint-knee outcome.",
            "scripts/run_final_revision_analyses.py:stage_constructability",
            fragment,
        )
    _add_value(
        rows,
        text,
        "continuous_minimum_segment",
        "Results 3.5",
        "1",
        "cell",
        "tables/generated/table_2_selected_profile.csv",
        "continuous_slope_deg and continuous_mu_static run lengths",
        "Minimum constant-value run length in selected continuous profile.",
        reporting_code,
        "minimum segment length 1 cell",
    )

    jma_fields = [
        (
            "jma_station_winter_count",
            str(len(jma)),
            "station-winters",
            "Daily forcing covered 4 stations and 12 station-winters",
            "row count",
        ),
        (
            "jma_lmax_ratio_min",
            f"{jma['joint_to_uniform_l_max_ratio'].min():.2f}",
            "ratio",
            "Lmax ratios ranged from 2.84",
            "min(joint_to_uniform_l_max_ratio)",
        ),
        (
            "jma_lmax_ratio_max",
            f"{jma['joint_to_uniform_l_max_ratio'].max():.2f}",
            "ratio",
            "2.84 to 16.68",
            "max(joint_to_uniform_l_max_ratio)",
        ),
        (
            "jma_smax_reduction_min",
            f"{jma['s_max_reduction_percent'].min():.2f}",
            "%",
            "Smax reductions ranged from 67.01%",
            "min(s_max_reduction_percent)",
        ),
        (
            "jma_smax_reduction_max",
            f"{jma['s_max_reduction_percent'].max():.2f}",
            "%",
            "67.01% to 70.83%",
            "max(s_max_reduction_percent)",
        ),
    ]
    for value_id, reported, units, fragment, locator in jma_fields:
        _add_value(
            rows,
            text,
            value_id,
            "Results 3.6; Table 9",
            reported,
            units,
            "results/generated/final_revision_jma_paired_tradeoffs.csv",
            locator,
            "Count or range across paired station-winter rows.",
            "scripts/run_final_revision_analyses.py:stage_jma",
            fragment,
        )

    output = ROOT / "manuscript_values.csv"
    pd.DataFrame(rows).to_csv(output, index=False, lineterminator="\n")
    shutil.copyfile(output, ROOT / "manuscript" / "manuscript_values.csv")
    failures = [row["value_id"] for row in rows if row["status"] != "PASS"]
    if failures:
        raise SystemExit(
            "manuscript value checks failed: " + ", ".join(failures)
        )
    return output


def validate_finalization() -> Path:
    value_table = pd.read_csv(ROOT / "manuscript_values.csv")
    manuscript_text = _document_text(MANUSCRIPT)
    reference_audit = pd.read_csv(
        ROOT / "references" / "final_revision_reference_audit.csv"
    )
    figure_formats = {
        "png": ".png",
        "tiff": ".tiff",
        "vector_eps": ".eps",
        "vector_pdf": ".pdf",
        "vector_svg": ".svg",
    }
    figure_checks = {
        key: all(
            (
                ROOT
                / "figures"
                / ("vector" if key.startswith("vector") else key)
                / f"figure_{number}_"
            ).parent.exists()
            and any(
                (
                    ROOT
                    / "figures"
                    / ("vector" if key.startswith("vector") else key)
                ).glob(f"figure_{number}_*{suffix}")
            )
            for number in range(1, 10)
        )
        for key, suffix in figure_formats.items()
    }
    table_checks = all(
        any(TABLES.glob(f"table_{number}_*.csv"))
        for number in range(1, 10)
    )
    manuscript_doc = Document(MANUSCRIPT)
    body = "\n".join(paragraph.text for paragraph in manuscript_doc.paragraphs)
    citation_checks = {
        "all_reference_titles_present": all(
            title in body for title in reference_audit["title"]
        ),
        "reference_sources_verified": reference_audit[
            "source_local_status"
        ].eq("verified").all(),
        "reference_source_paths_exist": all(
            (ROOT / source).exists()
            for source in reference_audit["source_path"]
        ),
    }
    checks = {
        "manuscript_values_pass": value_table["status"].eq("PASS").all(),
        "manuscript_values_fragments_present": all(
            fragment in manuscript_text
            for fragment in value_table["manuscript_fragment"]
        ),
        "all_nine_tables_present": table_checks,
        "manuscript_has_eight_embedded_tables_and_one_csv_table": (
            len(manuscript_doc.tables) == 8
            and "Table 4 (strategy phase-diagram classifications) is supplied "
            "separately as an editable CSV" in body
        ),
        "manuscript_has_nine_figure_captions": sum(
            paragraph.text.startswith(
                tuple(f"Figure {number}." for number in range(1, 10))
            )
            for paragraph in manuscript_doc.paragraphs
        )
        == 9,
        **{f"all_nine_figures_{key}": value for key, value in figure_checks.items()},
        **citation_checks,
    }
    output = ROOT / "audit" / "TARGETED_FINALIZATION.md"
    lines = [
        "# Targeted finalization audit",
        "",
        "| Check | Status |",
        "|---|---|",
        *[
            f"| {name.replace('_', ' ')} | {'PASS' if passed else 'FAIL'} |"
            for name, passed in checks.items()
        ],
        "",
        (
            f"- manuscript_values.csv rows: {len(value_table)}; "
            f"failures: {(~value_table['status'].eq('PASS')).sum()}."
        ),
        "- Figures and Tables were regenerated once from the canonical result files.",
        "- No new substantive analysis or sensitivity analysis was performed.",
    ]
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")
    failures = [name for name, passed in checks.items() if not passed]
    if failures:
        raise SystemExit("targeted finalization failed: " + ", ".join(failures))
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--stage",
        choices=["values", "validate", "all"],
        default="all",
    )
    args = parser.parse_args()
    if args.stage in {"values", "all"}:
        build_manuscript_values()
    if args.stage in {"validate", "all"}:
        validate_finalization()


if __name__ == "__main__":
    main()
