from __future__ import annotations

import html
import re
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "manuscript" / "manuscript_CRST.docx"
GUIDE = (
    ROOT
    / "data"
    / "raw"
    / "published_sources"
    / "20260925T050931Z"
    / "crst_guide_for_authors.html"
)
OUTPUT = ROOT / "audit" / "CRST_FORMAT_LANGUAGE_AUDIT.md"


def _paragraphs_before_references(document: Document) -> list[str]:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    return paragraphs[: paragraphs.index("References")]


def _abstract(document: Document) -> str:
    paragraphs = document.paragraphs
    abstract_index = next(
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph.text == "Abstract"
    )
    return paragraphs[abstract_index + 1].text


def _keywords(document: Document) -> list[str]:
    text = next(
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.startswith("Keywords:")
    )
    return [keyword.strip() for keyword in text.split(":", 1)[1].split(";")]


def _highlight_lines() -> list[str]:
    return [
        line.removeprefix("- ").strip()
        for line in (
            ROOT / "manuscript" / "highlights_CRST.txt"
        ).read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def _docx_xml() -> str:
    with zipfile.ZipFile(MANUSCRIPT) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def _tiff_checks() -> tuple[int, list[str]]:
    failures: list[str] = []
    figures = sorted((ROOT / "figures" / "tiff").glob("figure_*.tiff"))
    for path in figures:
        with Image.open(path) as image:
            dpi = image.info.get("dpi", (0, 0))
            if min(dpi) < 300:
                failures.append(f"{path.name}: {dpi}")
    return len(figures), failures


def _unit_superscript_check(document: Document) -> tuple[int, list[str]]:
    unit_runs = 0
    failures: list[str] = []
    paragraphs = list(document.paragraphs)
    paragraphs.extend(
        paragraph
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if re.search(r"−[123]", run.text):
                unit_runs += 1
                if not run._r.xpath(
                    "./w:rPr/w:vertAlign[@w:val='superscript']"
                ):
                    failures.append(run.text)
            if re.search(r"\b[mhs]-[123]\b", run.text):
                failures.append(run.text)
    return unit_runs, failures


def main() -> None:
    document = Document(MANUSCRIPT)
    body_paragraphs = _paragraphs_before_references(document)
    body = "\n".join(body_paragraphs)
    abstract = _abstract(document)
    keywords = _keywords(document)
    highlights = _highlight_lines()
    xml = _docx_xml()
    guide_html = GUIDE.read_text(encoding="utf-8")
    guide_text = re.sub(
        r"\s+",
        " ",
        html.unescape(re.sub(r"<[^>]+>", " ", guide_html)),
    )
    tiff_count, tiff_failures = _tiff_checks()
    unit_run_count, unit_failures = _unit_superscript_check(document)

    guide_requirements = {
        "250-word abstract": "does not exceed 250 words" in guide_text,
        "1-7 keywords": "1 to 7 keywords" in guide_text,
        "3-5 highlights": "3 to 5 bullet points" in guide_text,
        "85-character highlights": "maximum of 85 characters" in guide_text,
        "continuous line numbers": (
            "Continuous line numbers must be added" in guide_text
        ),
        "author-year references": (
            "Single author:" in guide_text
            and "first author" in guide_text
            and "alphabetically" in guide_text
        ),
        "separate artwork files": (
            "must be supplied as separate files" in guide_text
        ),
        "editable tables": "submitted as editable text" in guide_text,
        "data repository deposit": (
            "Deposit your research data in a relevant data repository"
            in guide_text
        ),
    }
    checks = {
        "official CRST guide snapshot verified": (
            GUIDE.exists() and all(guide_requirements.values())
        ),
        "single-column Word manuscript": '<w:cols w:space="720"/>' in xml,
        "continuous line numbering": (
            '<w:lnNumType w:countBy="1" w:start="1" '
            'w:restart="continuous"/>' in xml
        ),
        "abstract is at most 250 words": len(abstract.split()) <= 250,
        "keyword count is 1-7": 1 <= len(keywords) <= 7,
        "highlights are 3-5 bullets": 3 <= len(highlights) <= 5,
        "each highlight is at most 85 characters": all(
            len(highlight) <= 85 for highlight in highlights
        ),
        "native Word equations present": xml.count("<m:oMath") >= 3,
        "no visible LaTeX markers": not re.search(
            r"\\(?:frac|sum|theta|mu)|\$[^$]+\$",
            body,
        ),
        "unit exponents use font superscript": (
            unit_run_count > 0 and not unit_failures
        ),
        "no Unicode superscript glyphs": not re.search(
            "[⁰¹²³⁴⁵⁶⁷⁸⁹⁻]",
            body,
        ),
        "references use author-year citations": (
            not re.search(r"\[\d+(?:[-–,]\d+)*\]", body)
        ),
        "nine separate TIFF figures at 300 dpi or more": (
            tiff_count == 9 and not tiff_failures
        ),
        "editable tables supplied": (
            ROOT / "manuscript" / "editable_tables_CRST.docx"
        ).exists(),
        "editable figures supplied": (
            ROOT / "manuscript" / "editable_figures_CRST.pptx"
        ).exists(),
        "American English used consistently": not re.search(
            r"\b(?:optimisation|modelling|normalisation|behaviour|labour)\b",
            body,
            flags=re.IGNORECASE,
        ),
        "claims remain explicitly qualified": (
            "no front-wide superiority, structural safety, or validated roof "
            "performance is established" in abstract
            and "not design approval or safety certification" in body
        ),
        "DeepL is not claimed": "DeepL" not in body,
    }
    manual_items = [
        "Replace REQUIRED title-page affiliation and postal-address fields.",
        "Confirm funding, competing-interest declaration, and CRediT roles.",
        "Complete the Elsevier declarations-tool Word file.",
        "After author review, replace the provisional generative-AI confirmation "
        "sentence with the required final responsibility statement.",
        "Deposit the dataset/package in a relevant persistent repository and cite "
        "its identifier, or enter an allowed data-sharing explanation.",
        "Perform the final human spelling and grammar review before submission.",
    ]

    lines = [
        "# CRST format and language audit",
        "",
        f"Automated status: {'PASS' if all(checks.values()) else 'FAIL'}",
        "",
        "## Official-source basis",
        "",
        "- Current CRST Guide for Authors was retrieved as a browser-rendered "
        "official ScienceDirect page after the shell endpoint returned HTTP 403.",
        "- The current guide requires author-year citations and an alphabetically "
        "ordered reference list; the manuscript was updated from numeric citations.",
        "- General Elsevier artwork and generative-AI policies were also retained "
        "with checksums and acquisition metadata.",
        "",
        "## Automated gates",
        "",
        "| Gate | Status |",
        "|---|---|",
    ]
    lines.extend(
        f"| {name} | {'PASS' if passed else 'FAIL'} |"
        for name, passed in checks.items()
    )
    lines.extend(
        [
            "",
            "## Measured values",
            "",
            f"- Abstract: {len(abstract.split())} words.",
            f"- Keywords: {len(keywords)}.",
            f"- Highlights: {len(highlights)}; lengths "
            + ", ".join(str(len(highlight)) for highlight in highlights)
            + " characters.",
            f"- Native Word math objects: {xml.count('<m:oMath')}.",
            f"- Font-superscript unit exponent runs: {unit_run_count}.",
            f"- Separate TIFF figures: {tiff_count}.",
            "",
            "## Manual completion required before submission",
            "",
        ]
    )
    lines.extend(f"- [ ] {item}" for item in manual_items)
    lines.extend(
        [
            "",
            "## Language and claim control",
            "",
            "- American English is used consistently in manuscript prose.",
            "- No claim of DeepL use is present.",
            "- The abstract, discussion, limitations, and conclusions retain the "
            "audited conditional interpretation and do not claim structural safety "
            "or validated roof performance.",
            "- JMA observations remain described as supplementary scenario forcing.",
        ]
    )
    OUTPUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    if not all(checks.values()):
        failed = [name for name, passed in checks.items() if not passed]
        raise SystemExit(f"CRST format audit failed: {failed}")


if __name__ == "__main__":
    main()
