from docx import Document
from PIL import Image

from graded_roof.manuscript import _enforce_ascii
from graded_roof.validation import _reference_entries, _unsupported_non_ascii


def test_ascii_normalization_preserves_embedded_images(tmp_path):
    image_path = tmp_path / "figure.png"
    Image.new("RGB", (20, 20), "white").save(image_path)
    document = Document()
    document.add_paragraph("modeled–result")
    document.add_picture(str(image_path))

    _enforce_ascii(document)

    assert document.paragraphs[0].text == "modeled-result"
    assert len(document.inline_shapes) == 1


def test_ascii_normalization_formats_unit_exponents_as_superscript():
    document = Document()
    document.add_paragraph("80 kg m−1 and 2 W m−2")

    _enforce_ascii(document)

    paragraph = document.paragraphs[0]
    assert paragraph.text == "80 kg m−1 and 2 W m−2"
    superscript_runs = [
        run.text
        for run in paragraph.runs
        if run._r.xpath("./w:rPr/w:vertAlign[@w:val='superscript']")
    ]
    assert superscript_runs == ["−1", "−2"]


def test_validation_allows_unit_minus_but_rejects_other_non_ascii():
    assert _unsupported_non_ascii("80 kg m−1") == []
    assert _unsupported_non_ascii("smart quote ’") == ["’"]


def test_validation_counts_author_year_reference_entries():
    document = Document()
    document.add_paragraph("References")
    document.add_paragraph("Alpha, A., 2024. Example.")
    document.add_paragraph("Beta, B., 2025. Example.")
    document.add_paragraph("")
    document.add_paragraph("Editable tables")

    assert _reference_entries(document) == [
        "Alpha, A., 2024. Example.",
        "Beta, B., 2025. Example.",
    ]
